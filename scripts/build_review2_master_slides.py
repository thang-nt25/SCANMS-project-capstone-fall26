import os
import math
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml import parse_xml
from PIL import Image, ImageDraw, ImageFont
import fitz # PyMuPDF

W_INCHES = 13.333
H_INCHES = 7.5
W_PIXELS = 3000
H_PIXELS = 1688

PPTX_OUT = "docs/review 2/SCANMS_Thuyet_Trinh_Review_2_Va_Hoi_Dong.pptx"
PDF_OUT = "docs/review 2/SCANMS_Thuyet_Trinh_Review_2_Va_Hoi_Dong.pdf"
SLIDES_DIR = "scratch/generated_slides"
os.makedirs(SLIDES_DIR, exist_ok=True)
os.makedirs("docs/review 2", exist_ok=True)

def to_x(px):
    return Inches(W_INCHES * (px / float(W_PIXELS)))

def to_y(py):
    return Inches(H_INCHES * (py / float(H_PIXELS)))

def to_w(pw):
    return Inches(W_INCHES * (pw / float(W_PIXELS)))

def to_h(ph):
    return Inches(H_INCHES * (ph / float(H_PIXELS)))

# Color constants (Warm Sand & Brand Gold tokens)
C_INK = RGBColor(26, 22, 18)            # #1A1612
C_MUTED = RGBColor(125, 113, 94)        # #7D715E
C_BRAND = RGBColor(197, 155, 88)        # #C59B58
C_BRAND_STRONG = RGBColor(184, 142, 79) # #B88E4F
C_SAND_BG = RGBColor(243, 239, 230)     # #F3EFE6
C_SAND_BORDER = RGBColor(234, 228, 215) # #EAE4D7
C_WHITE = RGBColor(255, 255, 255)
C_GREEN = RGBColor(5, 150, 105)         # #059669
C_RED = RGBColor(220, 38, 38)           # #DC2626
C_DARK = RGBColor(35, 29, 21)           # #231D15
C_BRAND_SOFT = RGBColor(251, 245, 235)  # #FBF5EB

# PIL Colors
P_INK = (26, 22, 18)
P_MUTED = (125, 113, 94)
P_BRAND = (197, 155, 88)
P_BRAND_STRONG = (184, 142, 79)
P_SAND_BG = (243, 239, 230)
P_SAND_BORDER = (234, 228, 215)
P_WHITE = (255, 255, 255)
P_GREEN = (5, 150, 105)
P_RED = (220, 38, 38)
P_DARK = (35, 29, 21)
P_BRAND_SOFT = (251, 245, 235)

# PIL Fonts
try:
    font_hero = ImageFont.truetype("arialbd.ttf", 90)
    font_title = ImageFont.truetype("arialbd.ttf", 64)
    font_flow_title = ImageFont.truetype("arialbd.ttf", 44)
    font_sub = ImageFont.truetype("arialbd.ttf", 36)
    font_actor = ImageFont.truetype("arialbd.ttf", 30)
    font_label = ImageFont.truetype("arialbd.ttf", 22)
    font_small = ImageFont.truetype("arial.ttf", 16)
    font_content_num = ImageFont.truetype("arialbd.ttf", 46)
    font_content_txt = ImageFont.truetype("arialbd.ttf", 32)
    font_card_title = ImageFont.truetype("arialbd.ttf", 28)
    font_card_desc = ImageFont.truetype("arial.ttf", 19)
    font_badge = ImageFont.truetype("arialbd.ttf", 17)
    font_body_bold = ImageFont.truetype("arialbd.ttf", 20)
    font_body = ImageFont.truetype("arial.ttf", 19)
except Exception:
    font_hero = font_title = font_flow_title = font_sub = font_actor = font_label = font_small = font_content_num = font_content_txt = font_card_title = font_card_desc = font_badge = font_body_bold = font_body = ImageFont.load_default()

# ------------------------------------------------------------------------------
# PPTX HELPERS
# ------------------------------------------------------------------------------
def add_bg(slide, bg_path):
    slide.shapes.add_picture(bg_path, 0, 0, width=Inches(W_INCHES), height=Inches(H_INCHES))

def add_editable_text(slide, left_px, top_px, width_px, height_px, text, font_size_pt=14, bold=False, color=C_INK, align=PP_ALIGN.LEFT, font_name="Arial"):
    tx = slide.shapes.add_textbox(to_x(left_px), to_y(top_px), to_w(width_px), to_h(height_px))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = Pt(font_size_pt)
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = align
    return tx

def add_card(slide, left_px, top_px, width_px, height_px, fill_color=C_WHITE, border_color=C_SAND_BORDER, border_width_pt=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(left_px), to_y(top_px), to_w(width_px), to_h(height_px))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(border_width_pt)
    return shape

def add_image(slide, img_path, cx_px, cy_px, w_px, h_px):
    left_px = cx_px - w_px // 2
    top_px = cy_px - h_px // 2
    return slide.shapes.add_picture(img_path, to_x(left_px), to_y(top_px), to_w(w_px), to_h(h_px))

def add_arrow_line(slide, x0_px, y0_px, x1_px, y1_px, color=C_INK, width_pt=2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, to_x(x0_px), to_y(y0_px), to_x(x1_px), to_y(y1_px))
    conn.line.color.rgb = color
    conn.line.width = Pt(width_pt)
    line_pr = conn._element.spPr.ln
    arrow_xml = parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle" w="med" len="med"/>')
    line_pr.append(arrow_xml)
    return conn

def add_straight_line(slide, x0_px, y0_px, x1_px, y1_px, color=C_INK, width_pt=2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, to_x(x0_px), to_y(y0_px), to_x(x1_px), to_y(y1_px))
    conn.line.color.rgb = color
    conn.line.width = Pt(width_pt)
    return conn

# ------------------------------------------------------------------------------
# PIL HELPERS
# ------------------------------------------------------------------------------
def paste_icon(canvas, icon_path, center, size):
    im = Image.open(icon_path)
    im = im.resize(size, Image.Resampling.LANCZOS)
    cx, cy = center
    px = cx - size[0] // 2
    py = cy - size[1] // 2
    if im.mode == "RGBA":
        canvas.paste(im, (px, py), im)
    else:
        canvas.paste(im, (px, py))

def draw_pil_arrow(draw, start, end, label=None, label_pos=None, label_anchor="mm", width=4, head_len=18, color=(0,0,0)):
    x0, y0 = start
    x1, y1 = end
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(y1 - y0, x1 - x0)
    tip = (x1, y1)
    p1 = (x1 - head_len * math.cos(angle - math.pi / 6), y1 - head_len * math.sin(angle - math.pi / 6))
    p2 = (x1 - head_len * math.cos(angle + math.pi / 6), y1 - head_len * math.sin(angle + math.pi / 6))
    draw.polygon([tip, p1, p2], fill=color)
    if label:
        if label_pos is None:
            label_pos = ((x0 + x1) // 2, (y0 + y1) // 2 - 25)
        draw.text(label_pos, label, fill=color, font=font_label, anchor=label_anchor)

def wrap_text_pil(text, font, max_width, draw):
    lines = []
    for paragraph in text.split("\n"):
        words = paragraph.split(" ")
        cur_line = ""
        for w in words:
            test_line = (cur_line + " " + w).strip() if cur_line else w
            bbox = draw.textbbox((0, 0), test_line, font=font)
            w_px = bbox[2] - bbox[0]
            if w_px > max_width and cur_line:
                lines.append(cur_line)
                cur_line = w
            else:
                cur_line = test_line
        if cur_line:
            lines.append(cur_line)
    return lines

# ==============================================================================
# SLIDE 01: TITLE SLIDE (Bìa Thuyết Minh)
# ==============================================================================
def build_slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_title.png")

    # Native PPTX Textboxes
    add_editable_text(slide, 320, 480, 1400, 130, "SCANMS", font_size_pt=54, bold=True, color=C_INK)
    add_editable_text(slide, 320, 620, 1500, 140, "SALES COLLABORATOR & AFFILIATE\nNETWORK MANAGEMENT SYSTEM", font_size_pt=24, bold=True, color=C_BRAND_STRONG)
    add_straight_line(slide, 320, 800, 1300, 800, color=C_BRAND, width_pt=2.5)

    meta_text = (
        "CAPSTONE PROJECT DEFENSE  |  REVIEW 2\n"
        "Team: FA26SE032\n"
        "Team Leader: NGUYỄN THÀNH THẮNG\n"
        "Major: Software Engineering  |  FPT University"
    )
    add_editable_text(slide, 320, 840, 1500, 240, meta_text, font_size_pt=18, bold=False, color=C_INK)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_title.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((320, 520), "SCANMS", fill=P_INK, font=font_hero, anchor="lm")
    d.text((320, 660), "SALES COLLABORATOR & AFFILIATE\nNETWORK MANAGEMENT SYSTEM", fill=P_BRAND_STRONG, font=font_sub, anchor="lm")
    d.line([(320, 800), (1300, 800)], fill=P_BRAND, width=4)
    meta_y = 860
    d.text((320, meta_y), "CAPSTONE PROJECT DEFENSE  |  REVIEW 2", fill=P_BRAND_STRONG, font=font_actor, anchor="lm")
    d.text((320, meta_y + 60), "Team: FA26SE032", fill=P_MUTED, font=font_actor, anchor="lm")
    d.text((320, meta_y + 120), "Team Leader: NGUYỄN THÀNH THẮNG", fill=P_INK, font=font_actor, anchor="lm")
    d.text((320, meta_y + 180), "Major: Software Engineering  |  FPT University", fill=P_MUTED, font=font_actor, anchor="lm")
    bg.save(f"{SLIDES_DIR}/slide_01.png")
    print("Slide 01 generated.")

# ==============================================================================
# SLIDE 02: AGENDA (Mục Lục)
# ==============================================================================
def build_slide_02(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_content.png")

    items = [
        ("01", "OVERVIEW ARCHITECTURE (Kiến trúc Hệ thống Tổng quan)"),
        ("02", "CORE TECHNOLOGIES & APPARATUS (Công nghệ Cốt lõi & 4 Phân hệ Nền tảng)"),
        ("03", "CONCEPTUAL DATA MODEL (Mô hình Dữ liệu Thực thể Liên kết 16 Bảng)"),
        ("04", "5 MAIN BUSINESS FLOWS & UAT TEST MATRICES (5 Quy trình Nghiệp vụ & Kịch bản Test)"),
        ("05", "QUESTIONS & ANSWERS (Giải đáp Câu hỏi & Thảo luận Chuyên môn)")
    ]

    bg = Image.open("scratch/slide_assets/bg_content.png").convert("RGB")
    d = ImageDraw.Draw(bg)

    start_y = 440
    spacing = 225
    for i, (num, text) in enumerate(items):
        cy = start_y + i * spacing
        # PPTX Native Card
        add_card(slide, 320, cy, 2360, 175, fill_color=C_SAND_BG, border_color=C_SAND_BORDER, border_width_pt=1.5)
        add_editable_text(slide, 380, cy + 45, 120, 80, num, font_size_pt=34, bold=True, color=C_BRAND_STRONG)
        add_straight_line(slide, 520, cy + 30, 520, cy + 145, color=C_SAND_BORDER, width_pt=1.5)
        add_editable_text(slide, 560, cy + 50, 2060, 70, text, font_size_pt=20, bold=True, color=C_INK)

        # PIL Render
        d.rounded_rectangle([(320, cy), (2680, cy + 175)], radius=20, fill=P_SAND_BG, outline=P_SAND_BORDER, width=3)
        d.text((430, cy + 88), num, fill=P_BRAND_STRONG, font=font_content_num, anchor="mm")
        d.line([(520, cy + 30), (520, cy + 145)], fill=P_SAND_BORDER, width=3)
        d.text((560, cy + 88), text, fill=P_INK, font=font_content_txt, anchor="lm")

    bg.save(f"{SLIDES_DIR}/slide_02.png")
    print("Slide 02 generated.")

# ==============================================================================
# SLIDE 03: TECHNOLOGIES & 4 PORTALS (Công Nghệ & 4 Phân Hệ)
# ==============================================================================
def build_slide_03(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    # Left Brand Header
    add_editable_text(slide, 160, 500, 760, 80, "SCANMS", font_size_pt=46, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 160, 580, 760, 50, "MULTI-PORTAL PLATFORM", font_size_pt=18, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 160, 640, 760, 200, "Nền tảng tích hợp 4 cổng chuyên biệt kết nối chặt chẽ Khách hàng, KOL/KOC, Chủ Shop và Chuyên viên vận hành trên cùng một hạ tầng.", font_size_pt=14, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    apps = [
        ("CUSTOMER WEB APP", "Storefront Marketplace, Live Shopping, Strict Buyer Auth & VietQR Gateway\nCổng mua sắm mượt mà, xác thực người mua chặt chẽ, giỏ hàng bảo toàn và thanh toán tức thì.", "scratch/flow_icons/icon_cart.png"),
        ("COLLABORATOR PORTAL", "Affiliate Smart Links, Media Hub, Performance Analytics & Exclusive Deals\nCổng dành riêng cho KOL/KOC sinh Smart Link/QR, xin mẫu thử 14 ngày, đàm phán deal VIP và nhận hoa hồng.", "scratch/flow_icons/icon_link.png"),
        ("SHOP MANAGER PORTAL", "6-Block Product Matrix, Deal Negotiation, Order Fulfillment & GHN Delivery\nCổng vận hành Merchant: cấu hình ma trận sản phẩm 6 khối, duyệt hàng mẫu, mốc thưởng tháng (FR-09).", "scratch/flow_icons/icon_product_box.png"),
        ("SYSTEM ADMIN & MANAGER CONSOLE", "Platform Operations, Multi-Tier KYC, User RBAC Governance & Napas 24/7 Payouts\nTrung tâm điều hành: Chuyên viên thẩm định KYC/sản phẩm/trọng tài, Super Admin độc quyền quản trị và kho bạc.", "scratch/flow_icons/avatar_pristine.png")
    ]

    card_x = 1000
    start_y = 260
    card_w = 1840
    card_h = 270
    spacing = 310

    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)

    d.text((540, 530), "SCANMS", fill=P_INK, font=font_hero, anchor="mm")
    d.text((540, 610), "MULTI-PORTAL PLATFORM", fill=P_BRAND_STRONG, font=font_sub, anchor="mm")
    d.text((540, 670), "Nền tảng tích hợp 4 cổng chuyên biệt kết nối chặt chẽ Khách hàng, KOL/KOC, Chủ Shop", fill=P_MUTED, font=font_small, anchor="mm")
    d.text((540, 705), "và Chuyên viên vận hành trên cùng một hạ tầng.", fill=P_MUTED, font=font_small, anchor="mm")

    for i, (title, desc, icon_path) in enumerate(apps):
        cy = start_y + i * spacing
        # PPTX Native Card
        add_card(slide, card_x, cy, card_w, card_h, fill_color=C_WHITE, border_color=C_SAND_BORDER, border_width_pt=1.5)
        add_image(slide, icon_path, card_x + 130, cy + 135, 110, 110)
        add_editable_text(slide, card_x + 240, cy + 45, card_w - 280, 50, title, font_size_pt=18, bold=True, color=C_INK)
        add_editable_text(slide, card_x + 240, cy + 105, card_w - 280, 130, desc, font_size_pt=13, bold=False, color=C_MUTED)

        # PIL Render
        d.rounded_rectangle([(card_x, cy), (card_x + card_w, cy + card_h)], radius=24, fill=P_WHITE, outline=P_SAND_BORDER, width=3)
        paste_icon(bg, icon_path, (card_x + 130, cy + 135), (110, 110))
        d.text((card_x + 240, cy + 70), title, fill=P_INK, font=font_card_title, anchor="lm")
        lines = desc.split("\n")
        d.text((card_x + 240, cy + 125), lines[0], fill=P_MUTED, font=font_small, anchor="lm")
        if len(lines) > 1:
            d.text((card_x + 240, cy + 165), lines[1], fill=P_MUTED, font=font_small, anchor="lm")

    bg.save(f"{SLIDES_DIR}/slide_03.png")
    print("Slide 03 generated.")

# ==============================================================================
# SLIDE 04: ARCHITECTURE OVERVIEW (Kiến Trúc Hệ Thống Tổng Quan)
# ==============================================================================
def build_slide_04(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 85, 2200, 45, "SYSTEM ARCHITECTURAL OVERVIEW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 140, 2200, 60, "MULTI-TIER ARCHITECTURE WITH DISTRIBUTED SERVICES & INTEGRATIONS", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    arch_img_path = "docs/review 2/SCANMS_OVERVIEW_ARCHITECTURE.png"
    add_image(slide, arch_img_path, W_PIXELS // 2, 940, 2750, 1340)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 105), "SYSTEM ARCHITECTURAL OVERVIEW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 165), "MULTI-TIER ARCHITECTURE WITH DISTRIBUTED SERVICES & INTEGRATIONS", fill=P_INK, font=font_flow_title, anchor="mm")
    arch_im = Image.open(arch_img_path)
    arch_im.thumbnail((2750, 1340), Image.Resampling.LANCZOS)
    paste_x = (W_PIXELS - arch_im.width) // 2
    paste_y = 270 + (1340 - arch_im.height) // 2
    bg.paste(arch_im, (paste_x, paste_y))
    bg.save(f"{SLIDES_DIR}/slide_04.png")
    print("Slide 04 generated.")

# ==============================================================================
# SLIDE 05: CONCEPTUAL DATA MODEL (ERD)
# ==============================================================================
def build_slide_05(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 85, 2200, 45, "CONCEPTUAL DATA MODEL", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 140, 2200, 60, "CORE BUSINESS DOMAINS & CHEN NOTATION CARDINALITIES (16 CORE ENTITIES)", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    erd_img_path = "docs/review 2/SCANMS_ERD_01_Conceptual.png"
    add_image(slide, erd_img_path, W_PIXELS // 2, 940, 2750, 1340)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 105), "CONCEPTUAL DATA MODEL", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 165), "CORE BUSINESS DOMAINS & CHEN NOTATION CARDINALITIES (16 CORE ENTITIES)", fill=P_INK, font=font_flow_title, anchor="mm")
    erd_im = Image.open(erd_img_path)
    erd_im.thumbnail((2750, 1340), Image.Resampling.LANCZOS)
    paste_x = (W_PIXELS - erd_im.width) // 2
    paste_y = 270 + (1340 - erd_im.height) // 2
    bg.paste(erd_im, (paste_x, paste_y))
    bg.save(f"{SLIDES_DIR}/slide_05.png")
    print("Slide 05 generated.")

# ==============================================================================
# SLIDE 06: USER ACTORS & 5 ROLES
# ==============================================================================
def build_slide_06(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 85, 2200, 45, "SECURITY & ACCESS CONTROL MATRIX", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 140, 2200, 60, "SCANMS ACTORS & 5 STANDARDIZED RBAC ROLES", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    roles = [
        ("CUSTOMER", "Khách Mua Hàng", "Shopper & End Buyer", "Duyệt Marketplace, xem livestream, mua hàng qua link tiếp thị, theo dõi vận đơn GHN, khiếu nại đổi trả."),
        ("COLLABORATOR", "Đối Tác Tiếp Thị", "KOL / KOC Affiliate", "Định danh cá nhân, sinh Smart Link / Dynamic QR / Coupon độc quyền, xin mẫu thử 14 ngày, hưởng hoa hồng."),
        ("SHOP MANAGER", "Chủ Gian Hàng", "Merchant & Store Owner", "Quản lý sản phẩm 6 khối, duyệt hàng mẫu, deal VIP 2 chiều, cấu hình mốc thưởng tháng (FR-09), đối soát ví shop."),
        ("SYSTEM MANAGER", "Chuyên Viên Vận Hành", "Operations & Moderator", "Thẩm định hồ sơ KYC, kiểm duyệt sản phẩm mới, trọng tài hàng mẫu & phân xử khiếu nại đổi trả độc lập."),
        ("SYSTEM ADMIN", "Quản Trị Tối Cao", "Super Admin Governance", "Toàn quyền vận hành + ĐỘC QUYỀN: Quản trị tài khoản/phân quyền (/admin/users) & Báo cáo tài chính tổng (/admin/analytics).")
    ]

    card_w = 480
    card_h = 1180
    start_x = 160
    spacing = 540
    top_y = 260

    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 105), "SECURITY & ACCESS CONTROL MATRIX", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 165), "SCANMS ACTORS & 5 STANDARDIZED RBAC ROLES", fill=P_INK, font=font_flow_title, anchor="mm")

    for i, (name, vn_name, role, desc) in enumerate(roles):
        card_left = start_x + i * spacing
        cx = card_left + card_w // 2

        # PPTX Native Card
        add_card(slide, card_left, top_y, card_w, card_h, fill_color=C_WHITE, border_color=C_SAND_BORDER, border_width_pt=1.5)
        add_image(slide, "scratch/flow_icons/avatar_pristine.png", cx, top_y + 170, 240, 240)
        add_editable_text(slide, card_left + 20, top_y + 320, card_w - 40, 45, name, font_size_pt=18, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
        add_editable_text(slide, card_left + 20, top_y + 375, card_w - 40, 35, vn_name, font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
        add_editable_text(slide, card_left + 20, top_y + 415, card_w - 40, 35, role, font_size_pt=12, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)
        add_straight_line(slide, card_left + 40, top_y + 470, card_left + card_w - 40, top_y + 470, color=C_SAND_BORDER, width_pt=1)
        add_editable_text(slide, card_left + 30, top_y + 500, card_w - 60, 460, desc, font_size_pt=13, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

        # PIL Render
        d.rounded_rectangle([(card_left, top_y), (card_left + card_w, top_y + card_h)], radius=24, fill=P_WHITE, outline=P_SAND_BORDER, width=3)
        paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (cx, top_y + 170), (240, 240))
        d.text((cx, top_y + 340), name, fill=P_INK, font=font_card_title, anchor="mm")
        d.text((cx, top_y + 390), vn_name, fill=P_BRAND_STRONG, font=font_label, anchor="mm")
        d.text((cx, top_y + 430), role, fill=P_MUTED, font=font_small, anchor="mm")
        d.line([(card_left + 40, top_y + 470), (card_left + card_w - 40, top_y + 470)], fill=P_SAND_BORDER, width=2)
        words = desc.split(" ")
        lines, cur = [], ""
        for w in words:
            if len(cur + " " + w) > 28:
                lines.append(cur)
                cur = w
            else:
                cur = (cur + " " + w).strip()
        if cur:
            lines.append(cur)
        desc_y = top_y + 520
        for l in lines:
            d.text((cx, desc_y), l, fill=P_MUTED, font=font_small, anchor="mm")
            desc_y += 38

    bg.save(f"{SLIDES_DIR}/slide_06.png")
    print("Slide 06 generated.")

# ==============================================================================
# SLIDE 07: FLOW 1: USER REGISTER & KYC VERIFICATION (DIAGRAM)
# ==============================================================================
def build_slide_07(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 90, 2200, 45, "MAIN BUSINESS FLOW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 145, 2200, 60, "FLOW 1: USER REGISTRATION & MULTI-TIER KYC VERIFICATION", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Actor Left: COLLABORATOR / SHOP
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 320, 800, 220, 220)
    add_editable_text(slide, 120, 930, 400, 50, "COLLABORATOR / SHOP", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 1: Form checklist
    add_arrow_line(slide, 460, 800, 830, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 500, 755, 300, 40, "ENTER INFORMATION", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/form_checklist_clean.png", 940, 785, 200, 200)

    # Step 2: Phone OTP
    add_arrow_line(slide, 1050, 800, 1340, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 1080, 755, 250, 40, "VERIFY 2FA / OTP", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/phone_otp_clean.png", 1440, 785, 170, 190)

    # Step 3: ID Card
    add_arrow_line(slide, 1540, 800, 1890, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 1580, 755, 280, 40, "UPLOAD DOSSIER", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/id_card_clean.png", 2020, 790, 220, 140)
    add_editable_text(slide, 1820, 890, 400, 40, "CCCD + STK / MST + LICENSE", font_size_pt=11, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Step 4: Folder Upload
    add_arrow_line(slide, 2140, 800, 2410, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 2180, 755, 220, 40, "SUBMIT KYC", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/folder_clean.png", 2520, 800, 160, 180)

    # Actor Top Right: SYSTEM MANAGER
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 2520, 340, 220, 220)
    add_editable_text(slide, 2320, 460, 400, 45, "SYSTEM MANAGER", font_size_pt=16, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_arrow_line(slide, 2520, 520, 2520, 700, color=C_INK, width_pt=2)
    add_editable_text(slide, 2540, 580, 280, 40, "AUDIT KYC DOSSIER", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.LEFT)

    # Branch 1: APPROVE
    add_straight_line(slide, 2610, 800, 2750, 800, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2750, 800, 2750, 1220, color=C_INK, width_pt=2)
    add_editable_text(slide, 2760, 880, 180, 40, "APPROVE", font_size_pt=13, bold=True, color=C_GREEN, align=PP_ALIGN.LEFT)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 2750, 1340, 180, 180)
    add_editable_text(slide, 2550, 1460, 400, 50, "ACTIVE ROLE ACCOUNT", font_size_pt=16, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    # Branch 2: REJECT
    add_straight_line(slide, 2480, 890, 2480, 1220, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2480, 1220, 2150, 1220, color=C_INK, width_pt=2)
    add_editable_text(slide, 2220, 1175, 180, 40, "REJECT", font_size_pt=13, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_reject_x.png", 2020, 1220, 170, 170)
    add_arrow_line(slide, 2020, 1120, 2020, 930, color=C_INK, width_pt=2)
    add_editable_text(slide, 1680, 1020, 320, 40, "RE-SUBMIT WITH REASON", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.RIGHT)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 110), "MAIN BUSINESS FLOW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 175), "FLOW 1: USER REGISTRATION & MULTI-TIER KYC VERIFICATION", fill=P_INK, font=font_flow_title, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (320, 800), (220, 220))
    d.text((320, 950), "COLLABORATOR / SHOP", fill=P_INK, font=font_actor, anchor="mm")
    draw_pil_arrow(d, (460, 800), (830, 800), "ENTER INFORMATION", (645, 760))
    paste_icon(bg, "scratch/flow_icons/form_checklist_clean.png", (940, 785), (200, 200))
    draw_pil_arrow(d, (1050, 800), (1340, 800), "VERIFY 2FA / OTP", (1195, 760))
    paste_icon(bg, "scratch/flow_icons/phone_otp_clean.png", (1440, 785), (170, 190))
    draw_pil_arrow(d, (1540, 800), (1890, 800), "UPLOAD DOSSIER", (1715, 760))
    paste_icon(bg, "scratch/flow_icons/id_card_clean.png", (2020, 790), (220, 140))
    d.text((2020, 900), "CCCD + STK / MST + LICENSE", fill=P_MUTED, font=font_badge, anchor="mm")
    draw_pil_arrow(d, (2140, 800), (2410, 800), "SUBMIT KYC", (2275, 760))
    paste_icon(bg, "scratch/flow_icons/folder_clean.png", (2520, 800), (160, 180))
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (2520, 340), (220, 220))
    d.text((2240, 340), "SYSTEM MANAGER", fill=P_INK, font=font_actor, anchor="rm")
    draw_pil_arrow(d, (2520, 480), (2520, 700), "AUDIT KYC DOSSIER", (2660, 590), label_anchor="lm")
    d.line([(2610, 800), (2750, 800), (2750, 1220)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2750, 1220), (2750, 1240), head_len=18)
    d.text((2770, 880), "APPROVE", fill=P_GREEN, font=font_label, anchor="lm")
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (2750, 1340), (180, 180))
    d.text((2750, 1470), "ACTIVE ROLE ACCOUNT", fill=P_GREEN, font=font_actor, anchor="mm")
    d.line([(2480, 890), (2480, 1220), (2160, 1220)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2160, 1220), (2130, 1220), head_len=18)
    d.text((2320, 1180), "REJECT", fill=P_RED, font=font_label, anchor="mm")
    paste_icon(bg, "scratch/extracted_elements/icon_reject_x.png", (2020, 1220), (170, 170))
    draw_pil_arrow(d, (2020, 1120), (2020, 930), "RE-SUBMIT WITH REASON", (1850, 1020), label_anchor="rm")
    bg.save(f"{SLIDES_DIR}/slide_07.png")
    print("Slide 07 generated.")

# ==============================================================================
# REUSABLE HELPER: COMPANION TEST MATRIX SLIDE (2 BALANCED CARDS)
# ==============================================================================
def build_test_matrix_slide(prs, flow_num, flow_title_en, cases_succ, cases_fail, slide_idx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    # Header PPTX
    add_editable_text(slide, 600, 85, 2200, 45, "MAIN BUSINESS FLOW - TEST & VERIFICATION MATRIX", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 140, 2200, 60, f"FLOW {flow_num}: {flow_title_en} - UAT TEST CASES", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Card dimensions
    card_w = 1280
    card_h = 1260
    top_y = 260
    left_x1 = 180
    left_x2 = 1540

    # -----------------------------
    # PPTX Left Card: Happy Path
    # -----------------------------
    add_card(slide, left_x1, top_y, card_w, card_h, fill_color=C_WHITE, border_color=C_SAND_BORDER, border_width_pt=1.5)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", left_x1 + card_w // 2, top_y + 110, 140, 140)
    add_editable_text(slide, left_x1 + 60, top_y + 195, card_w - 120, 40, "KỲ VỌNG THÀNH CÔNG (HAPPY PATH)", font_size_pt=18, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, left_x1 + 60, top_y + 240, card_w - 120, 35, "Kết quả trả về khi dữ liệu hợp lệ & luồng xử lý hoàn tất", font_size_pt=13, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_straight_line(slide, left_x1 + 60, top_y + 285, left_x1 + card_w - 60, top_y + 285, color=C_SAND_BORDER, width_pt=1.5)

    cur_y = top_y + 320
    for head, desc in cases_succ:
        add_card(slide, left_x1 + 60, cur_y, card_w - 120, 185, fill_color=C_SAND_BG, border_color=C_SAND_BORDER, border_width_pt=1.5)
        # Dot indicator
        add_card(slide, left_x1 + 90, cur_y + 25, 18, 18, fill_color=C_GREEN, border_color=C_GREEN)
        add_editable_text(slide, left_x1 + 125, cur_y + 20, card_w - 200, 35, head, font_size_pt=13, bold=True, color=C_INK)
        add_editable_text(slide, left_x1 + 125, cur_y + 60, card_w - 200, 110, desc, font_size_pt=12, bold=False, color=C_MUTED)
        cur_y += 205

    # Bottom status tag
    add_card(slide, left_x1 + 60, top_y + card_h - 75, card_w - 120, 50, fill_color=C_BRAND_SOFT, border_color=C_SAND_BORDER, border_width_pt=1)
    add_editable_text(slide, left_x1 + 60, top_y + card_h - 65, card_w - 120, 35, "KẾT QUẢ THỰC NGHIỆM: 4/4 TEST SCENARIOS ĐẠT CHUẨN (100% PASS)", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # -----------------------------
    # PPTX Right Card: Fail-Safe
    # -----------------------------
    add_card(slide, left_x2, top_y, card_w, card_h, fill_color=C_WHITE, border_color=C_SAND_BORDER, border_width_pt=1.5)
    add_image(slide, "scratch/extracted_elements/icon_reject_x.png", left_x2 + card_w // 2, top_y + 110, 140, 140)
    add_editable_text(slide, left_x2 + 60, top_y + 195, card_w - 120, 40, "XỬ LÝ THẤT BẠI & NGOẠI LỆ (FAIL-SAFE)", font_size_pt=18, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, left_x2 + 60, top_y + 240, card_w - 120, 35, "Cơ chế phòng vệ, mã lỗi HTTP và phản hồi an toàn cho người dùng", font_size_pt=13, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_straight_line(slide, left_x2 + 60, top_y + 285, left_x2 + card_w - 60, top_y + 285, color=C_SAND_BORDER, width_pt=1.5)

    cur_y = top_y + 320
    for head, desc in cases_fail:
        add_card(slide, left_x2 + 60, cur_y, card_w - 120, 185, fill_color=C_SAND_BG, border_color=C_SAND_BORDER, border_width_pt=1.5)
        # Dot indicator
        add_card(slide, left_x2 + 90, cur_y + 25, 18, 18, fill_color=C_RED, border_color=C_RED)
        add_editable_text(slide, left_x2 + 125, cur_y + 20, card_w - 200, 35, head, font_size_pt=13, bold=True, color=C_INK)
        add_editable_text(slide, left_x2 + 125, cur_y + 60, card_w - 200, 110, desc, font_size_pt=12, bold=False, color=C_MUTED)
        cur_y += 205

    # Bottom status tag
    add_card(slide, left_x2 + 60, top_y + card_h - 75, card_w - 120, 50, fill_color=C_BRAND_SOFT, border_color=C_SAND_BORDER, border_width_pt=1)
    add_editable_text(slide, left_x2 + 60, top_y + card_h - 65, card_w - 120, 35, "CƠ CHẾ PHÒNG VỆ: BẢO MẬT ĐA TẦNG & TUYỆT ĐỐI KHÔNG NUỐT LỖI", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # Footer note PPTX
    add_editable_text(slide, 200, 1570, 2600, 35, "Tài liệu chi tiết 20 Test Cases UAT: docs/review 2/SCANMS_Demo_Test_Cases_Review2.xlsx", font_size_pt=11, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    # -----------------------------
    # PIL Render
    # -----------------------------
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 100), "MAIN BUSINESS FLOW - TEST & VERIFICATION MATRIX", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 160), f"FLOW {flow_num}: {flow_title_en} - UAT TEST CASES", fill=P_INK, font=font_flow_title, anchor="mm")

    # PIL Left Card
    d.rounded_rectangle([(left_x1, top_y), (left_x1 + card_w, top_y + card_h)], radius=24, fill=P_WHITE, outline=P_SAND_BORDER, width=3)
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (left_x1 + card_w // 2, top_y + 110), (140, 140))
    d.text((left_x1 + card_w // 2, top_y + 210), "KỲ VỌNG THÀNH CÔNG (HAPPY PATH)", fill=P_INK, font=font_card_title, anchor="mm")
    d.text((left_x1 + card_w // 2, top_y + 250), "Kết quả trả về khi dữ liệu hợp lệ & luồng xử lý hoàn tất", fill=P_BRAND_STRONG, font=font_sub, anchor="mm")
    d.line([(left_x1 + 60, top_y + 285), (left_x1 + card_w - 60, top_y + 285)], fill=P_SAND_BORDER, width=2)

    cur_y = top_y + 320
    for head, desc in cases_succ:
        d.rounded_rectangle([(left_x1 + 60, cur_y), (left_x1 + card_w - 60, cur_y + 185)], radius=14, fill=P_SAND_BG, outline=P_SAND_BORDER, width=2)
        d.ellipse([(left_x1 + 90, cur_y + 24), (left_x1 + 106, cur_y + 40)], fill=P_GREEN)
        d.text((left_x1 + 125, cur_y + 32), head, fill=P_INK, font=font_body_bold, anchor="lm")
        lines = wrap_text_pil(desc, font_body, card_w - 180, d)
        ly = cur_y + 75
        for l in lines:
            d.text((left_x1 + 125, ly), l, fill=P_MUTED, font=font_body, anchor="lm")
            ly += 30
        cur_y += 205

    d.rounded_rectangle([(left_x1 + 60, top_y + card_h - 75), (left_x1 + card_w - 60, top_y + card_h - 25)], radius=10, fill=P_BRAND_SOFT, outline=P_SAND_BORDER, width=1)
    d.text((left_x1 + card_w // 2, top_y + card_h - 50), "KẾT QUẢ THỰC NGHIỆM: 4/4 TEST SCENARIOS ĐẠT CHUẨN (100% PASS)", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")

    # PIL Right Card
    d.rounded_rectangle([(left_x2, top_y), (left_x2 + card_w, top_y + card_h)], radius=24, fill=P_WHITE, outline=P_SAND_BORDER, width=3)
    paste_icon(bg, "scratch/extracted_elements/icon_reject_x.png", (left_x2 + card_w // 2, top_y + 110), (140, 140))
    d.text((left_x2 + card_w // 2, top_y + 210), "XỬ LÝ THẤT BẠI & NGOẠI LỆ (FAIL-SAFE)", fill=P_INK, font=font_card_title, anchor="mm")
    d.text((left_x2 + card_w // 2, top_y + 250), "Cơ chế phòng vệ, mã lỗi HTTP và phản hồi an toàn cho người dùng", fill=P_RED, font=font_sub, anchor="mm")
    d.line([(left_x2 + 60, top_y + 285), (left_x2 + card_w - 60, top_y + 285)], fill=P_SAND_BORDER, width=2)

    cur_y = top_y + 320
    for head, desc in cases_fail:
        d.rounded_rectangle([(left_x2 + 60, cur_y), (left_x2 + card_w - 60, cur_y + 185)], radius=14, fill=P_SAND_BG, outline=P_SAND_BORDER, width=2)
        d.ellipse([(left_x2 + 90, cur_y + 24), (left_x2 + 106, cur_y + 40)], fill=P_RED)
        d.text((left_x2 + 125, cur_y + 32), head, fill=P_INK, font=font_body_bold, anchor="lm")
        lines = wrap_text_pil(desc, font_body, card_w - 180, d)
        ly = cur_y + 75
        for l in lines:
            d.text((left_x2 + 125, ly), l, fill=P_MUTED, font=font_body, anchor="lm")
            ly += 30
        cur_y += 205

    d.rounded_rectangle([(left_x2 + 60, top_y + card_h - 75), (left_x2 + card_w - 60, top_y + card_h - 25)], radius=10, fill=P_BRAND_SOFT, outline=P_SAND_BORDER, width=1)
    d.text((left_x2 + card_w // 2, top_y + card_h - 50), "CƠ CHẾ PHÒNG VỆ: BẢO MẬT ĐA TẦNG & TUYỆT ĐỐI KHÔNG NUỐT LỖI", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")

    d.text((1500, 1580), "Tài liệu chi tiết 20 Test Cases UAT: docs/review 2/SCANMS_Demo_Test_Cases_Review2.xlsx", fill=P_MUTED, font=font_small, anchor="mm")

    out_img = f"{SLIDES_DIR}/slide_{slide_idx:02d}.png"
    bg.save(out_img)
    print(f"Slide {slide_idx:02d} (Flow {flow_num} Test Matrix) generated.")

# ==============================================================================
# SLIDE 08: FLOW 1 TEST MATRIX
# ==============================================================================
def build_slide_08(prs):
    succ = [
        ("TC_KYC_01: Đăng ký tài khoản (FR-01, FR-02)", "HTTP 201 Created; Mật khẩu mã hóa Argon2id an toàn.\nTài khoản ở trạng thái UNVERIFIED; Điều hướng sang trang /login."),
        ("TC_KYC_02: Đăng nhập an toàn & 2FA OTP (FR-03)", "HTTP 200 OK; Cấp JWT Access & Refresh Token trong HttpOnly Cookie.\nRBAC Guard đối soát vai trò, tự động điều hướng đúng Portal."),
        ("TC_KYC_03: Nộp hồ sơ định danh KYC (FR-05)", "Regex kiểm tra CCCD (12 số) & STK; Upload ảnh Cloudinary CDN.\nHồ sơ chuyển PENDING_REVIEW; Hiển thị huy hiệu 'Chờ thẩm định'."),
        ("TC_KYC_04: System Manager phê duyệt KYC (FR-06)", "Hồ sơ chuyển sang APPROVED; Kích hoạt toàn bộ tính năng tiếp thị.\nKOL được phép sinh Smart Link, tạo mã QR và hưởng hoa hồng.")
    ]
    fail = [
        ("TC_KYC_01: Trùng Email / SĐT hoặc Mật khẩu yếu", "Trùng lặp: HTTP 409 Conflict, thông báo đỏ 'Email/SĐT đã tồn tại'.\nMật khẩu yếu: HTTP 422 Unprocessable, báo lỗi chi tiết tại ô nhập."),
        ("TC_KYC_02: Sai thông tin đăng nhập hoặc Quá lượt OTP", "Sai pass: HTTP 401 Unauthorized kèm thông báo bảo mật.\nSai OTP quá 5 lần: HTTP 429 Rate Limit, khóa tạm 15 phút chống brute-force."),
        ("TC_KYC_03: Ảnh sai định dạng hoặc CCCD không đúng chuẩn", "File >5MB hoặc sai định dạng: Chặn tải lên, báo 'Chỉ nhận JPG/PNG < 5MB'.\nCCCD sai 12 số: Form báo đỏ tại chỗ, vô hiệu hóa submit."),
        ("TC_KYC_04: Hồ sơ KYC mờ / Không đạt yêu cầu", "System Manager chọn REJECT bắt buộc nhập lý do; Hồ sơ chuyển REJECTED.\nGửi email yêu cầu nộp lại ảnh rõ nét; Tài khoản giữ nguyên.")
    ]
    build_test_matrix_slide(prs, 1, "USER REGISTRATION & MULTI-TIER KYC", succ, fail, 8)

# ==============================================================================
# SLIDE 09: FLOW 2: 2-WAY COMMISSION NEGOTIATION & LIVESTREAM ENGINE (DIAGRAM)
# ==============================================================================
def build_slide_09(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 90, 2200, 45, "MAIN BUSINESS FLOW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 145, 2200, 60, "FLOW 2: 2-WAY COMMISSION DEAL NEGOTIATION & LIVESTREAM ENGINE", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Actor 1: COLLABORATOR (Top Left)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 300, 500, 200, 200)
    add_editable_text(slide, 120, 620, 360, 45, "COLLABORATOR", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 120, 665, 360, 35, "(KOL / KOC Affiliate)", font_size_pt=12, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Actor 2: SHOP MANAGER (Bottom Left)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 300, 1100, 200, 200)
    add_editable_text(slide, 120, 1220, 360, 45, "SHOP MANAGER", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 120, 1265, 360, 35, "(Merchant Store Owner)", font_size_pt=12, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Step 1: Base Floor Rate
    add_arrow_line(slide, 420, 1100, 690, 1100, color=C_INK, width_pt=2)
    add_editable_text(slide, 430, 1055, 250, 40, "PUBLISH PRODUCT", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_product_box.png", 790, 1100, 160, 160)
    add_editable_text(slide, 630, 1195, 320, 40, "BASE 10% FLOOR RATE\n(Standard Open Offer)", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # Discover Product wire
    add_straight_line(slide, 790, 1000, 790, 740, color=C_INK, width_pt=2)
    add_straight_line(slide, 790, 740, 510, 740, color=C_INK, width_pt=2)
    add_arrow_line(slide, 510, 740, 510, 640, color=C_INK, width_pt=2)
    add_editable_text(slide, 520, 700, 260, 35, "DISCOVER PRODUCT", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 2: 2-Way Negotiation
    add_arrow_line(slide, 420, 500, 860, 500, color=C_INK, width_pt=2)
    add_editable_text(slide, 440, 455, 400, 40, "PROPOSE 18% VIP RATE", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_handshake.png", 960, 500, 160, 160)
    add_editable_text(slide, 800, 605, 320, 45, "PROPOSAL & COMMITMENT\n(Video Quota & GMV Target)", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    add_straight_line(slide, 890, 1100, 1210, 1100, color=C_INK, width_pt=2)
    add_straight_line(slide, 1210, 1100, 1210, 840, color=C_INK, width_pt=2)
    add_arrow_line(slide, 1210, 840, 1310, 840, color=C_INK, width_pt=2)
    add_editable_text(slide, 940, 1055, 260, 40, "SHOP DEAL INVITE\n(Chat Socket.IO)", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Handshake into Contract
    add_straight_line(slide, 1050, 500, 1210, 500, color=C_INK, width_pt=2)
    add_straight_line(slide, 1210, 500, 1210, 760, color=C_INK, width_pt=2)
    add_arrow_line(slide, 1210, 760, 1310, 760, color=C_INK, width_pt=2)
    add_editable_text(slide, 1070, 455, 140, 40, "DEAL AGREED", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 3: Exclusive Contract
    add_image(slide, "scratch/flow_icons/icon_deal_contract.png", 1420, 800, 180, 190)
    add_editable_text(slide, 1220, 640, 400, 45, "EXCLUSIVE DEAL CONTRACT", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 1220, 685, 400, 35, "(FR-27 2-WAY NEGOTIATION)", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # Step 4: Decisions
    add_straight_line(slide, 1520, 840, 1620, 840, color=C_INK, width_pt=2)
    add_straight_line(slide, 1620, 840, 1620, 1140, color=C_INK, width_pt=2)
    add_arrow_line(slide, 1620, 1140, 1860, 1140, color=C_INK, width_pt=2)
    add_editable_text(slide, 1660, 1095, 180, 40, "REJECT DEAL", font_size_pt=12, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_reject_x.png", 1960, 1140, 150, 150)
    add_editable_text(slide, 1780, 1235, 360, 50, "FALLBACK TO FLOOR RATE\nMaintains 10% Base Open Offer", font_size_pt=12, bold=True, color=C_RED, align=PP_ALIGN.CENTER)

    add_arrow_line(slide, 1520, 760, 1860, 760, color=C_INK, width_pt=2)
    add_editable_text(slide, 1560, 715, 260, 40, "APPROVE VIP DEAL", font_size_pt=12, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 1960, 760, 160, 160)
    add_editable_text(slide, 1780, 650, 360, 45, "DEAL ACCEPTED (18% VIP)", font_size_pt=15, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    # Step 5: Multi-Rule Boosts
    add_arrow_line(slide, 2060, 760, 2240, 760, color=C_INK, width_pt=2)
    add_editable_text(slide, 2070, 715, 160, 40, "AUTO-STACK", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_tier_stack.png", 2340, 760, 150, 150)
    add_editable_text(slide, 2180, 865, 320, 50, "TIER RANK BOOST\n(Silver +1% | Gold +2% | Diamond +3%)", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    add_arrow_line(slide, 2440, 760, 2600, 760, color=C_INK, width_pt=2)
    add_editable_text(slide, 2440, 715, 160, 40, "+ LIVE BONUS", font_size_pt=11, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_livestream.png", 2700, 760, 150, 150)
    add_editable_text(slide, 2540, 865, 320, 50, "+2% TO +3% LIVE BONUS\n(Session Status = LIVE)", font_size_pt=11, bold=True, color=C_RED, align=PP_ALIGN.CENTER)

    # Step 6: Assets Issued
    add_straight_line(slide, 2340, 930, 2340, 1050, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2340, 1050, 2460, 1050, color=C_INK, width_pt=2)
    add_editable_text(slide, 2350, 1005, 140, 40, "ACTIVATE", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.LEFT)
    add_image(slide, "scratch/flow_icons/icon_link.png", 2560, 1050, 140, 140)
    add_editable_text(slide, 2400, 1140, 360, 65, "EXCLUSIVE ASSETS ISSUED\n• Custom Voucher Code (FR-12)\n• Free Sample Dispatched (FR-14)", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 110), "MAIN BUSINESS FLOW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 175), "FLOW 2: 2-WAY COMMISSION DEAL NEGOTIATION & LIVESTREAM ENGINE", fill=P_INK, font=font_flow_title, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (300, 500), (200, 200))
    d.text((300, 630), "COLLABORATOR", fill=P_INK, font=font_actor, anchor="mm")
    d.text((300, 670), "(KOL / KOC Affiliate)", fill=P_MUTED, font=font_small, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (300, 1100), (200, 200))
    d.text((300, 1230), "SHOP MANAGER", fill=P_INK, font=font_actor, anchor="mm")
    d.text((300, 1270), "(Merchant Store Owner)", fill=P_MUTED, font=font_small, anchor="mm")
    draw_pil_arrow(d, (420, 1100), (690, 1100), "PUBLISH PRODUCT", (555, 1060))
    paste_icon(bg, "scratch/flow_icons/icon_product_box.png", (790, 1100), (160, 160))
    d.text((790, 1205), "BASE 10% FLOOR RATE", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")
    d.text((790, 1235), "(Standard Open Offer)", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(790, 1000), (790, 740), (510, 740), (510, 640)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (510, 650), (510, 630), "DISCOVER PRODUCT", (650, 705), label_anchor="mm", head_len=18)
    draw_pil_arrow(d, (420, 500), (860, 500), "PROPOSE 18% VIP RATE", (640, 460))
    paste_icon(bg, "scratch/flow_icons/icon_handshake.png", (960, 500), (160, 160))
    d.text((960, 605), "PROPOSAL & COMMITMENT", fill=P_INK, font=font_badge, anchor="mm")
    d.text((960, 635), "(Video Quota & GMV Target)", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(1050, 500), (1210, 500), (1210, 760)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (1210, 760), (1310, 760), head_len=18)
    d.text((1130, 460), "DEAL AGREED", fill=P_INK, font=font_label, anchor="mm")
    d.line([(890, 1100), (1210, 1100), (1210, 840)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (1210, 840), (1310, 840), head_len=18)
    d.text((1050, 1060), "SHOP DEAL INVITE (Socket.IO)", fill=P_INK, font=font_label, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/icon_deal_contract.png", (1420, 800), (180, 190))
    d.text((1420, 650), "EXCLUSIVE DEAL CONTRACT", fill=P_INK, font=font_card_title, anchor="mm")
    d.text((1420, 685), "(FR-27 2-WAY NEGOTIATION)", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")
    d.line([(1520, 840), (1620, 840), (1620, 1140), (1860, 1140)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (1850, 1140), (1880, 1140), "REJECT DEAL", (1750, 1100), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_reject_x.png", (1960, 1140), (150, 150))
    d.text((1960, 1245), "FALLBACK TO FLOOR RATE", fill=P_RED, font=font_actor, anchor="mm")
    d.text((1960, 1285), "Maintains 10% Base Open Offer", fill=P_MUTED, font=font_small, anchor="mm")
    draw_pil_arrow(d, (1520, 760), (1860, 760), "APPROVE VIP DEAL", (1690, 720), head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (1960, 760), (160, 160))
    d.text((1960, 660), "DEAL ACCEPTED (18% VIP)", fill=P_GREEN, font=font_actor, anchor="mm")
    draw_pil_arrow(d, (2060, 760), (2240, 760), "AUTO-STACK", (2150, 720), head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_tier_stack.png", (2340, 760), (150, 150))
    d.text((2340, 875), "TIER RANK BOOST", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")
    d.text((2340, 905), "(Silver +1% | Gold +2% | Diamond +3%)", fill=P_MUTED, font=font_badge, anchor="mm")
    draw_pil_arrow(d, (2440, 760), (2600, 760), "+ LIVE BONUS", (2520, 720), head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_livestream.png", (2700, 760), (150, 150))
    d.text((2700, 875), "+2% TO +3% LIVE BONUS", fill=P_RED, font=font_badge, anchor="mm")
    d.text((2700, 905), "(Session Status = LIVE)", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(2340, 930), (2340, 1050), (2460, 1050)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2450, 1050), (2480, 1050), "ACTIVATE", (2380, 1010), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_link.png", (2560, 1050), (140, 140))
    d.text((2560, 1150), "EXCLUSIVE ASSETS ISSUED", fill=P_INK, font=font_badge, anchor="mm")
    d.text((2560, 1180), "• Custom Voucher Code (FR-12)", fill=P_MUTED, font=font_badge, anchor="mm")
    d.text((2560, 1210), "• Free Sample Dispatched (FR-14)", fill=P_MUTED, font=font_badge, anchor="mm")
    bg.save(f"{SLIDES_DIR}/slide_09.png")
    print("Slide 09 (Flow 2 Diagram) generated.")

# ==============================================================================
# SLIDE 10: FLOW 2 TEST MATRIX
# ==============================================================================
def build_slide_10(prs):
    succ = [
        ("TC_DEAL_01: Shop gửi lời mời Deal qua Chat (FR-27)", "Socket.IO phát event real-time, hiển thị thẻ Deal trên chat 2 bên.\nBản ghi DealContract khởi tạo với status = PENDING kèm cam kết GMV."),
        ("TC_DEAL_02: KOL chấp thuận Deal VIP & Sinh Voucher (FR-27)", "Deal chuyển status = ACTIVE; Thẻ chat đổi sang xanh 'ĐÃ CHẤP THUẬN'.\nTự động sinh Voucher Code độc quyền và kích hoạt rate VIP 18%."),
        ("TC_DEAL_03: Mở phòng Live & Kích hoạt Live Bonus +3% (FR-15)", "Phiên chuyển status = LIVE, mở Room Socket.IO nhận tương tác.\nEngine tự động cộng dồn Live Bonus (+3%), nâng tổng rate lên 21%."),
        ("TC_DEAL_04: Kết thúc phiên Live & Khóa tương tác (FR-15)", "Chuyển status = ENDED; Màn hình Ended Stream Overlay xuất hiện.\nKhóa chat/ghim deal; Gỡ thưởng Live Bonus đưa hoa hồng về 18%.")
    ]
    fail = [
        ("TC_DEAL_01: Đặt mức hoa hồng vượt trần sàn (>25%)", "Chặn ngay tại API, trả về HTTP 400: 'Hoa hồng không được vượt 25%'.\nMất kết nối Socket: Client tự động retry và hiển thị cảnh báo offline."),
        ("TC_DEAL_02: KOL từ chối Deal hoặc Đề xuất quá hạn 48h", "KOL từ chối: Deal chuyển REJECTED, sản phẩm duy trì rate sàn 10%.\nQuá hạn 48h: Thẻ tự động chuyển EXPIRED, khóa nút chấp thuận."),
        ("TC_DEAL_03: Sản phẩm hết hàng hoặc Hoa hồng vượt trần sàn", "Kho hết hàng: Giỏ Live tự động ẩn nút mua, báo 'Tạm hết hàng'.\nRate cộng dồn vượt 25%: Tự động cắt về trần 25% (Math.min(25, rate))."),
        ("TC_DEAL_04: Tắt tab đột ngột hoặc Nhắn tin sau khi đóng", "KOL tắt tab: Sau 5 phút timeout server tự động đóng phòng an toàn.\nKhán giả gửi tin sau live: Socket từ chối báo 'Phiên đã kết thúc'.")
    ]
    build_test_matrix_slide(prs, 2, "2-WAY COMMISSION DEAL & LIVESTREAM", succ, fail, 10)

# ==============================================================================
# SLIDE 11: FLOW 3: MULTI-CHANNEL AFFILIATE PROMOTION & CHECKOUT (DIAGRAM)
# ==============================================================================
def build_slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 90, 2200, 45, "MAIN BUSINESS FLOW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 145, 2200, 60, "FLOW 3: MULTI-CHANNEL AFFILIATE PROMOTION & BUYER CHECKOUT", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Actor 1: COLLABORATOR (Top Left)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 300, 480, 200, 200)
    add_editable_text(slide, 120, 600, 360, 45, "COLLABORATOR", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 120, 645, 360, 35, "(KOL / KOC Affiliate)", font_size_pt=12, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Actor 2: CUSTOMER (Bottom Left)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 300, 1120, 200, 200)
    add_editable_text(slide, 120, 1240, 360, 45, "CUSTOMER", font_size_pt=15, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 120, 1285, 360, 35, "(Marketplace Shopper)", font_size_pt=12, bold=False, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Channel 1: Smart Link & QR
    add_arrow_line(slide, 420, 480, 620, 480, color=C_INK, width_pt=2)
    add_editable_text(slide, 430, 435, 200, 40, "CREATE ATTRIBUTION", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_link.png", 710, 480, 150, 150)
    add_editable_text(slide, 550, 570, 320, 45, "SMART LINK & QR (FR-10, 11)\n30-Day Last-Click Cookie", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Channel 2: Livestream Shopping
    add_straight_line(slide, 420, 520, 440, 520, color=C_INK, width_pt=2)
    add_straight_line(slide, 440, 520, 440, 780, color=C_INK, width_pt=2)
    add_arrow_line(slide, 440, 780, 620, 780, color=C_INK, width_pt=2)
    add_editable_text(slide, 450, 735, 170, 40, "HOST LIVE (FR-15)", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_livestream.png", 710, 780, 160, 160)
    add_editable_text(slide, 550, 875, 320, 45, "LIVESTREAM SHOPPING\nReal-time Deal Pin (+3% Live Bonus)", font_size_pt=11, bold=True, color=C_RED, align=PP_ALIGN.CENTER)

    # Convergence into Cart
    add_straight_line(slide, 800, 480, 1060, 480, color=C_INK, width_pt=2)
    add_straight_line(slide, 1060, 480, 1060, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 820, 435, 220, 40, "CLICK SMART LINK", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    add_straight_line(slide, 800, 780, 1060, 780, color=C_INK, width_pt=2)
    add_editable_text(slide, 820, 735, 220, 40, "PIN LIVE DEALS", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    add_straight_line(slide, 420, 1120, 1060, 1120, color=C_INK, width_pt=2)
    add_straight_line(slide, 1060, 1120, 1060, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 600, 1075, 300, 40, "DISCOVER & BROWSE", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    add_arrow_line(slide, 1060, 880, 1260, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 1080, 835, 170, 40, "ADD TO CART", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_cart.png", 1360, 880, 160, 160)
    add_editable_text(slide, 1210, 980, 300, 40, "MARKETPLACE CART", font_size_pt=13, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 3: Strict Buyer Auth
    add_arrow_line(slide, 1450, 880, 1630, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 1460, 835, 170, 40, "SECURE CHECKOUT", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_lock_auth.png", 1720, 880, 150, 150)
    add_editable_text(slide, 1550, 980, 340, 50, "STRICT BUYER AUTH\nPopup Login / 100% Cart Preserved", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # Step 4: VietQR / PayOS
    add_arrow_line(slide, 1800, 880, 1950, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 1810, 835, 140, 40, "PAY VIETQR", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_qr_payment.png", 2040, 880, 140, 150)
    add_editable_text(slide, 1890, 980, 300, 45, "VIETQR / PAYOS / COD\nInstant Napas 24/7 Gateway", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 5: Webhook
    add_arrow_line(slide, 2120, 880, 2340, 880, color=C_INK, width_pt=2)
    add_editable_text(slide, 2130, 835, 200, 40, "WEBHOOK VERIFY", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_system_monitor.png", 2430, 880, 150, 150)
    add_editable_text(slide, 2280, 770, 300, 35, "PAYMENT GATEWAY", font_size_pt=12, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Branches
    add_straight_line(slide, 2510, 840, 2600, 840, color=C_INK, width_pt=2)
    add_straight_line(slide, 2600, 840, 2600, 700, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2600, 700, 2700, 700, color=C_INK, width_pt=2)
    add_editable_text(slide, 2550, 655, 180, 40, "PAID", font_size_pt=12, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 2780, 700, 140, 140)
    add_editable_text(slide, 2630, 785, 300, 45, "ORDER PAID\nPending Escrow Commission", font_size_pt=11, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    add_straight_line(slide, 2510, 920, 2600, 920, color=C_INK, width_pt=2)
    add_straight_line(slide, 2600, 920, 2600, 1060, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2600, 1060, 2700, 1060, color=C_INK, width_pt=2)
    add_editable_text(slide, 2550, 1015, 180, 40, "PENDING", font_size_pt=12, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_timer_escrow.png", 2780, 1060, 140, 140)
    add_editable_text(slide, 2630, 1145, 300, 45, "UNPAID ORDERS TAB\nPending Payment & Timeout", font_size_pt=11, bold=True, color=C_RED, align=PP_ALIGN.CENTER)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 110), "MAIN BUSINESS FLOW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 175), "FLOW 3: MULTI-CHANNEL AFFILIATE PROMOTION & BUYER CHECKOUT", fill=P_INK, font=font_flow_title, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (300, 480), (200, 200))
    d.text((300, 610), "COLLABORATOR", fill=P_INK, font=font_actor, anchor="mm")
    d.text((300, 650), "(KOL / KOC Affiliate)", fill=P_MUTED, font=font_small, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (300, 1120), (200, 200))
    d.text((300, 1250), "CUSTOMER", fill=P_INK, font=font_actor, anchor="mm")
    d.text((300, 1290), "(Marketplace Shopper)", fill=P_MUTED, font=font_small, anchor="mm")
    draw_pil_arrow(d, (420, 480), (620, 480), "CREATE ATTRIBUTION", (520, 440))
    paste_icon(bg, "scratch/flow_icons/icon_link.png", (710, 480), (150, 150))
    d.text((710, 580), "SMART LINK & QR (FR-10, 11)", fill=P_INK, font=font_badge, anchor="mm")
    d.text((710, 610), "30-Day Last-Click Cookie", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(420, 520), (440, 520), (440, 780), (620, 780)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (610, 780), (630, 780), head_len=18)
    d.text((530, 740), "HOST LIVE (FR-15)", fill=P_INK, font=font_label, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/icon_livestream.png", (710, 780), (160, 160))
    d.text((710, 885), "LIVESTREAM SHOPPING", fill=P_RED, font=font_badge, anchor="mm")
    d.text((710, 915), "Real-time Deal Pin (+3% Live Bonus)", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(800, 480), (1060, 480), (1060, 880)], fill=(0, 0, 0), width=4)
    d.text((930, 440), "CLICK SMART LINK", fill=P_INK, font=font_label, anchor="mm")
    d.line([(800, 780), (1060, 780)], fill=(0, 0, 0), width=4)
    d.text((930, 740), "PIN LIVE DEALS", fill=P_INK, font=font_label, anchor="mm")
    d.line([(420, 1120), (1060, 1120), (1060, 880)], fill=(0, 0, 0), width=4)
    d.text((740, 1080), "DISCOVER & BROWSE", fill=P_INK, font=font_label, anchor="mm")
    draw_pil_arrow(d, (1060, 880), (1260, 880), "ADD TO CART", (1160, 840))
    paste_icon(bg, "scratch/flow_icons/icon_cart.png", (1360, 880), (160, 160))
    d.text((1360, 990), "MARKETPLACE CART", fill=P_INK, font=font_actor, anchor="mm")
    draw_pil_arrow(d, (1450, 880), (1630, 880), "SECURE CHECKOUT", (1540, 840))
    paste_icon(bg, "scratch/flow_icons/icon_lock_auth.png", (1720, 880), (150, 150))
    d.text((1720, 990), "STRICT BUYER AUTH", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")
    d.text((1720, 1020), "Popup Login / 100% Cart Preserved", fill=P_MUTED, font=font_badge, anchor="mm")
    draw_pil_arrow(d, (1800, 880), (1950, 880), "PAY VIETQR", (1875, 840))
    paste_icon(bg, "scratch/flow_icons/icon_qr_payment.png", (2040, 880), (140, 150))
    d.text((2040, 990), "VIETQR / PAYOS / COD", fill=P_INK, font=font_badge, anchor="mm")
    d.text((2040, 1020), "Instant Napas 24/7 Gateway", fill=P_MUTED, font=font_badge, anchor="mm")
    draw_pil_arrow(d, (2120, 880), (2340, 880), "WEBHOOK VERIFY", (2230, 840))
    paste_icon(bg, "scratch/extracted_elements/icon_system_monitor.png", (2430, 880), (150, 150))
    d.text((2430, 780), "PAYMENT GATEWAY", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(2510, 840), (2600, 840), (2600, 700), (2700, 700)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2690, 700), (2720, 700), "PAID", (2640, 660), head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (2780, 700), (140, 140))
    d.text((2780, 795), "ORDER PAID", fill=P_GREEN, font=font_badge, anchor="mm")
    d.text((2780, 825), "Pending Escrow Commission", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(2510, 920), (2600, 920), (2600, 1060), (2700, 1060)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2690, 1060), (2720, 1060), "PENDING", (2640, 1020), head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_timer_escrow.png", (2780, 1060), (140, 140))
    d.text((2780, 1155), "UNPAID ORDERS TAB", fill=P_RED, font=font_badge, anchor="mm")
    d.text((2780, 1185), "Pending Payment & Timeout", fill=P_MUTED, font=font_badge, anchor="mm")
    bg.save(f"{SLIDES_DIR}/slide_11.png")
    print("Slide 11 (Flow 3 Diagram) generated.")

# ==============================================================================
# SLIDE 12: FLOW 3 TEST MATRIX
# ==============================================================================
def build_slide_12(prs):
    succ = [
        ("TC_ATTR_01: KOL sinh Smart Link & Dynamic QR (FR-10, 11)", "Sinh link rút gọn /r/:slug tích hợp sẵn collaboratorId và UTM tag.\nRender mã QR Vector sắc nét; Cho phép tải ảnh PNG chất lượng cao."),
        ("TC_ATTR_02: Khách bấm link & Ghi nhận Cookie 30 ngày (FR-13)", "Redirect 302 về trang sản phẩm; Tự động ghi Cookie scanms_ref 30 ngày.\nCơ chế Last-Click Attribution ghi đè mã người giới thiệu mới nhất."),
        ("TC_ATTR_03: Đăng nhập nhanh tại chỗ bảo toàn giỏ hàng (FR-16)", "Modal Quick Login mở tại chỗ; Đăng nhập thành công giữ 100% giỏ hàng.\nKhách không bị văng trang, chuyển ngay sang thanh toán (Strict Auth)."),
        ("TC_ATTR_04: Thanh toán VietQR PayOS & Webhook (FR-16, 17)", "Backend xác thực chữ ký HMAC-SHA256, chuyển đơn sang status = PAID.\nVí hoa hồng KOL ghi nhận phần thưởng ở trạng thái PENDING_ESCROW.")
    ]
    fail = [
        ("TC_ATTR_01: KOL chưa duyệt KYC hoặc Sản phẩm ngừng bán", "Chưa duyệt KYC: Nút tạo link bị khóa kèm cảnh báo 'Cần duyệt KYC'.\nSản phẩm đã xóa: Trả về HTTP 404 'Sản phẩm tạm dừng kinh doanh'."),
        ("TC_ATTR_02: Link hỏng hoặc Cookie đã quá hạn 30 ngày", "Link sai slug: Chuyển hướng về Marketplace, báo 'Link không tồn tại'.\nQuá hạn 30 ngày: Cookie tự hủy, đơn tính là đơn tự nhiên (0% hoa hồng)."),
        ("TC_ATTR_03: Khách chưa có tài khoản hoặc Nhập sai pass", "Chưa có tài khoản: Nút 'Đăng ký nhanh 30s' ngay trên Modal tiện lợi.\nSai mật khẩu: Báo đỏ trực tiếp trên Modal, không reload giỏ hàng."),
        ("TC_ATTR_04: Sai chữ ký Webhook hoặc Khách chuyển thiếu tiền", "Sai chữ ký Webhook: Backend từ chối với HTTP 403, ghi log bảo mật.\nChuyển thiếu tiền: Đơn giữ PENDING_PAYMENT, yêu cầu chuyển đủ.")
    ]
    build_test_matrix_slide(prs, 3, "AFFILIATE ATTRIBUTION & BUYER CHECKOUT", succ, fail, 12)

# ==============================================================================
# SLIDE 13: FLOW 4: FULFILLMENT, 14-DAY ESCROW & DISPUTES (DIAGRAM)
# ==============================================================================
def build_slide_13(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 90, 2200, 45, "MAIN BUSINESS FLOW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 145, 2200, 60, "FLOW 4: FULFILLMENT, 14-DAY ESCROW & DISPUTE ARBITRATION", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Actor 1: SHOP MANAGER (Left)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 320, 800, 220, 220)
    add_editable_text(slide, 120, 930, 400, 50, "SHOP MANAGER", font_size_pt=16, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Fulfillment Arrow
    add_arrow_line(slide, 460, 800, 820, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 500, 755, 280, 40, "CREATE GHN SHIPMENT", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_truck.png", 940, 800, 200, 180)

    # Deliver to Customer
    add_arrow_line(slide, 1060, 800, 1360, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 1100, 755, 220, 40, "DELIVER PARCEL", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Actor 2: CUSTOMER (Middle)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 1480, 800, 220, 220)
    add_editable_text(slide, 1280, 930, 400, 50, "CUSTOMER", font_size_pt=16, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Happy Path (Top Branch)
    add_straight_line(slide, 1610, 750, 1690, 750, color=C_INK, width_pt=2)
    add_straight_line(slide, 1690, 750, 1690, 560, color=C_INK, width_pt=2)
    add_arrow_line(slide, 1690, 560, 1980, 560, color=C_INK, width_pt=2)
    add_editable_text(slide, 1720, 515, 240, 40, "HOLD 14-DAY ESCROW", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_timer_escrow.png", 2090, 560, 170, 170)

    add_arrow_line(slide, 2190, 560, 2520, 560, color=C_INK, width_pt=2)
    add_editable_text(slide, 2230, 515, 240, 40, "14-DAY MATURED", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 2640, 560, 170, 170)
    add_editable_text(slide, 2420, 680, 440, 55, "ORDER COMPLETED & RELEASE FUNDS\nShop Receives Revenue, KOL Receives Commission", font_size_pt=13, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    # Dispute Path (Bottom Branch)
    add_straight_line(slide, 1610, 850, 1690, 850, color=C_INK, width_pt=2)
    add_straight_line(slide, 1690, 850, 1690, 1080, color=C_INK, width_pt=2)
    add_arrow_line(slide, 1690, 1080, 1940, 1080, color=C_INK, width_pt=2)
    add_editable_text(slide, 1730, 1035, 180, 40, "OPEN DISPUTE", font_size_pt=12, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_video_dispute.png", 2040, 1080, 160, 160)
    add_editable_text(slide, 1870, 1180, 340, 45, "BUYER UNBOXING VIDEO\n& RETURN REQUEST", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Sub-branch 1: Shop Agrees
    add_straight_line(slide, 2130, 1040, 2260, 1040, color=C_INK, width_pt=2)
    add_straight_line(slide, 2260, 1040, 2260, 960, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2260, 960, 2480, 960, color=C_INK, width_pt=2)
    add_editable_text(slide, 2270, 915, 210, 40, "SHOP ACCEPTS RETURN", font_size_pt=11, bold=True, color=C_GREEN, align=PP_ALIGN.LEFT)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 2580, 960, 140, 140)
    add_editable_text(slide, 2410, 1045, 340, 50, "REFUND BUYER & CLAWBACK\nCommission reversed from wallet", font_size_pt=11, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    # Sub-branch 2: Shop Rejects -> System Manager Arbitration
    add_straight_line(slide, 2130, 1120, 2260, 1120, color=C_INK, width_pt=2)
    add_straight_line(slide, 2260, 1120, 2260, 1260, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2260, 1260, 2480, 1260, color=C_INK, width_pt=2)
    add_editable_text(slide, 2270, 1215, 200, 40, "SHOP REJECTS", font_size_pt=11, bold=True, color=C_RED, align=PP_ALIGN.LEFT)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 2580, 1260, 140, 140)
    add_editable_text(slide, 2370, 1345, 420, 60, "SYSTEM MANAGER ARBITRATION\nAudits Unboxing Video vs Packing Camera\nFinal Binding Settlement & Escrow Decision", font_size_pt=11, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 110), "MAIN BUSINESS FLOW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 175), "FLOW 4: FULFILLMENT, 14-DAY ESCROW & DISPUTE ARBITRATION", fill=P_INK, font=font_flow_title, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (320, 800), (220, 220))
    d.text((320, 950), "SHOP MANAGER", fill=P_INK, font=font_actor, anchor="mm")
    draw_pil_arrow(d, (460, 800), (820, 800), "CREATE GHN SHIPMENT", (640, 760))
    paste_icon(bg, "scratch/flow_icons/icon_truck.png", (940, 800), (200, 180))
    draw_pil_arrow(d, (1060, 800), (1360, 800), "DELIVER PARCEL", (1210, 760))
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (1480, 800), (220, 220))
    d.text((1480, 950), "CUSTOMER", fill=P_INK, font=font_actor, anchor="mm")
    d.line([(1610, 750), (1690, 750), (1690, 560), (1980, 560)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (1960, 560), (1990, 560), "HOLD 14-DAY ESCROW", (1840, 520), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_timer_escrow.png", (2090, 560), (170, 170))
    draw_pil_arrow(d, (2190, 560), (2520, 560), "14-DAY MATURED", (2355, 520))
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (2640, 560), (170, 170))
    d.text((2640, 680), "ORDER COMPLETED & RELEASE FUNDS", fill=P_GREEN, font=font_actor, anchor="mm")
    d.text((2640, 720), "Shop Receives Revenue, KOL Receives Commission", fill=P_MUTED, font=font_small, anchor="mm")
    d.line([(1610, 850), (1690, 850), (1690, 1080), (1940, 1080)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (1930, 1080), (1960, 1080), "OPEN DISPUTE", (1825, 1040), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/flow_icons/icon_video_dispute.png", (2040, 1080), (160, 160))
    d.text((2040, 1185), "BUYER UNBOXING VIDEO", fill=P_INK, font=font_badge, anchor="mm")
    d.text((2040, 1215), "& RETURN REQUEST", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(2130, 1040), (2260, 1040), (2260, 960), (2480, 960)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2470, 960), (2500, 960), "SHOP ACCEPTS RETURN", (2370, 920), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (2580, 960), (140, 140))
    d.text((2580, 1045), "REFUND BUYER & CLAWBACK", fill=P_GREEN, font=font_badge, anchor="mm")
    d.text((2580, 1075), "Commission reversed from wallet", fill=P_MUTED, font=font_badge, anchor="mm")
    d.line([(2130, 1120), (2260, 1120), (2260, 1260), (2480, 1260)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2470, 1260), (2500, 1260), "SHOP REJECTS", (2370, 1220), label_anchor="mm", head_len=18)
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (2580, 1260), (140, 140))
    d.text((2580, 1350), "SYSTEM MANAGER ARBITRATION", fill=P_BRAND_STRONG, font=font_badge, anchor="mm")
    d.text((2580, 1380), "Audits Unboxing Video vs Packing Camera", fill=P_MUTED, font=font_badge, anchor="mm")
    d.text((2580, 1410), "Final Binding Settlement & Escrow Decision", fill=P_MUTED, font=font_badge, anchor="mm")
    bg.save(f"{SLIDES_DIR}/slide_13.png")
    print("Slide 13 (Flow 4 Diagram) generated.")

# ==============================================================================
# SLIDE 14: FLOW 4 TEST MATRIX
# ==============================================================================
def build_slide_14(prs):
    succ = [
        ("TC_FUL_01: Shop đóng gói & Tạo vận đơn GHN (FR-20)", "Gọi API GHN thành công; Sinh mã vận đơn, chuyển status = SHIPPING.\nKhởi tạo Timeline theo dõi hành trình giao nhận cho Khách và Shop."),
        ("TC_FUL_02: Khách nhận hàng & Khởi động Escrow 14 ngày (FR-21)", "Đơn cập nhật DELIVERED; Bắt đầu đếm ngược thời gian Escrow 14 ngày.\nTiền giữ tại tài khoản trung gian của Sàn; Hoa hồng KOL ở PENDING_ESCROW."),
        ("TC_FUL_03: Khách mở tranh chấp với Video Unboxing (FR-22)", "Đơn chuyển sang RETURN_REQUESTED; Đồng hồ Escrow LẬP TỨC ĐÓNG BĂNG.\nHồ sơ khiếu nại hiển thị trong tab 'Trả hàng / Hoàn tiền' của Shop."),
        ("TC_FUL_04: System Manager trọng tài đối soát 2 Video (FR-23)", "System Manager đối soát video 2 bên, ra phán quyết hoàn tiền 100% cho khách.\nCơ chế Clawback tự động: Thu hồi toàn bộ hoa hồng tạm khỏi ví KOL.")
    ]
    fail = [
        ("TC_FUL_01: Thiếu thông số đóng gói hoặc Cổng GHN bảo trì", "Thiếu cân nặng/kích thước: Form chặn submit, yêu cầu nhập số gram.\nCổng GHN gián đoạn: Báo 'Cổng vận chuyển bận, vui lòng thử lại', giữ đơn."),
        ("TC_FUL_02: Giao hàng thất bại 3 lần hoặc KOL cố rút tiền non", "Giao thất bại: Chuyển sang RETURNING (chuyển hoàn kho Shop), hủy hoa hồng.\nRút tiền non: Từ chối vì số dư khả dụng sau Escrow = 0đ, chống gian lận."),
        ("TC_FUL_03: Khiếu nại quá hạn 14 ngày hoặc Không có video", "Quá hạn 14 ngày: Nút khiếu nại bị ẩn, báo 'Đã hết thời hạn đổi trả'.\nKhông có video: Bắt buộc đính kèm video mở hộp mới cho phép gửi yêu cầu."),
        ("TC_FUL_04: Phán quyết Bác bỏ khiếu nại hoặc Shop im lặng", "Bác bỏ khiếu nại: Mở đóng băng Escrow, tiếp tục giải ngân bình thường.\nShop không phản hồi sau 48h: Tự động xử thắng cho Khách hàng.")
    ]
    build_test_matrix_slide(prs, 4, "FULFILLMENT & ESCROW ARBITRATION", succ, fail, 14)

# ==============================================================================
# SLIDE 15: FLOW 5: COMMISSION PAYOUT & NAPAS 24/7 TRANSFER (DIAGRAM)
# ==============================================================================
def build_slide_15(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_master.png")

    add_editable_text(slide, 600, 90, 2200, 45, "MAIN BUSINESS FLOW", font_size_pt=14, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 600, 145, 2200, 60, "FLOW 5: COMMISSION SETTLEMENT & NAPAS 24/7 TRANSFER", font_size_pt=20, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 1: System Trigger
    add_image(slide, "scratch/extracted_elements/icon_system_monitor.png", 320, 540, 180, 180)
    add_editable_text(slide, 120, 650, 400, 40, "AUTOMATED CRONJOB", font_size_pt=13, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    add_arrow_line(slide, 320, 700, 320, 830, color=C_INK, width_pt=2)
    add_editable_text(slide, 340, 750, 260, 40, "MATURE ESCROW FUNDS", font_size_pt=11, bold=True, color=C_INK, align=PP_ALIGN.LEFT)

    # Actor: COLLABORATOR
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 320, 960, 200, 200)
    add_editable_text(slide, 120, 1090, 400, 45, "COLLABORATOR", font_size_pt=16, bold=True, color=C_INK, align=PP_ALIGN.CENTER)

    # Step 2: Request Withdrawal
    add_arrow_line(slide, 440, 960, 830, 960, color=C_INK, width_pt=2)
    add_editable_text(slide, 480, 915, 300, 40, "REQUEST WITHDRAWAL", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_clipboard.png", 940, 960, 180, 190)
    add_editable_text(slide, 740, 835, 400, 40, "PAYOUT DOSSIER & 10% PIT TAX", font_size_pt=11, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Step 3: CRITICAL CORRECTION -> SYSTEM ADMIN (Super Admin)
    add_image(slide, "scratch/flow_icons/avatar_pristine.png", 940, 1340, 190, 190)
    add_editable_text(slide, 690, 1450, 500, 45, "SYSTEM ADMIN", font_size_pt=16, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 690, 1495, 500, 40, "(Super Admin Treasury Vault Governance)", font_size_pt=12, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)

    add_arrow_line(slide, 940, 1230, 940, 1080, color=C_INK, width_pt=2)
    add_editable_text(slide, 960, 1140, 360, 50, "AUDIT & APPROVE PAYOUT\n(EXCLUSIVELY SYSTEM ADMIN)", font_size_pt=12, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.LEFT)

    # Step 4: Dispatch Napas
    add_arrow_line(slide, 1050, 960, 1440, 960, color=C_INK, width_pt=2)
    add_editable_text(slide, 1090, 915, 280, 40, "DISPATCH NAPAS 24/7", font_size_pt=12, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/flow_icons/icon_cash_clean.png", 1560, 960, 200, 200)
    add_editable_text(slide, 1380, 1080, 360, 45, "NAPAS 24/7 GATEWAY\nInstant Bank Transfer API", font_size_pt=11, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    # Step 5: Settlement Results
    # Success
    add_straight_line(slide, 1680, 960, 2020, 960, color=C_INK, width_pt=2)
    add_straight_line(slide, 2020, 960, 2020, 800, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2020, 800, 2280, 800, color=C_INK, width_pt=2)
    add_editable_text(slide, 2060, 755, 180, 40, "SUCCESS", font_size_pt=12, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_approve_check.png", 2400, 800, 170, 170)
    add_editable_text(slide, 2200, 910, 400, 50, "BANK ACCOUNT CREDITED\nInstant 24/7 Transfer", font_size_pt=14, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)

    # Failed
    add_straight_line(slide, 2020, 960, 2020, 1180, color=C_INK, width_pt=2)
    add_arrow_line(slide, 2020, 1180, 2280, 1180, color=C_INK, width_pt=2)
    add_editable_text(slide, 2060, 1135, 180, 40, "FAILED", font_size_pt=12, bold=True, color=C_RED, align=PP_ALIGN.CENTER)
    add_image(slide, "scratch/extracted_elements/icon_reject_x.png", 2400, 1180, 160, 160)
    add_editable_text(slide, 2180, 1290, 440, 50, "RESTORE WALLET\nPessimistic Lock & Rollback", font_size_pt=14, bold=True, color=C_RED, align=PP_ALIGN.CENTER)

    # PIL Render
    bg = Image.open("scratch/slide_assets/bg_master.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((1700, 110), "MAIN BUSINESS FLOW", fill=P_BRAND_STRONG, font=font_label, anchor="mm")
    d.text((1700, 175), "FLOW 5: COMMISSION SETTLEMENT & NAPAS 24/7 TRANSFER", fill=P_INK, font=font_flow_title, anchor="mm")
    paste_icon(bg, "scratch/extracted_elements/icon_system_monitor.png", (320, 540), (180, 180))
    d.text((320, 660), "AUTOMATED CRONJOB", fill=P_MUTED, font=font_small, anchor="mm")
    draw_pil_arrow(d, (320, 700), (320, 830), "MATURE ESCROW FUNDS", (480, 765), label_anchor="lm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (320, 960), (200, 200))
    d.text((320, 1090), "COLLABORATOR", fill=P_INK, font=font_actor, anchor="mm")
    draw_pil_arrow(d, (440, 960), (830, 960), "REQUEST WITHDRAWAL", (635, 920))
    paste_icon(bg, "scratch/extracted_elements/icon_clipboard.png", (940, 960), (180, 190))
    d.text((940, 845), "PAYOUT DOSSIER & 10% PIT TAX", fill=P_MUTED, font=font_small, anchor="mm")
    paste_icon(bg, "scratch/flow_icons/avatar_pristine.png", (940, 1340), (190, 190))
    d.text((940, 1465), "SYSTEM ADMIN", fill=P_INK, font=font_actor, anchor="mm")
    d.text((940, 1505), "(Super Admin Treasury Vault Governance)", fill=P_BRAND_STRONG, font=font_small, anchor="mm")
    draw_pil_arrow(d, (940, 1230), (940, 1080), "AUDIT & APPROVE PAYOUT (EXCLUSIVELY SYSTEM ADMIN)", (970, 1155), label_anchor="lm", color=P_BRAND_STRONG)
    draw_pil_arrow(d, (1050, 960), (1440, 960), "DISPATCH NAPAS 24/7", (1245, 920))
    paste_icon(bg, "scratch/flow_icons/icon_cash_clean.png", (1560, 960), (200, 200))
    d.text((1560, 1085), "NAPAS 24/7 GATEWAY", fill=P_MUTED, font=font_small, anchor="mm")
    d.text((1560, 1115), "Instant Bank Transfer API", fill=P_MUTED, font=font_small, anchor="mm")
    d.line([(1680, 960), (2020, 960), (2020, 800), (2280, 800)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2270, 800), (2300, 800), "SUCCESS", (2150, 760), head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_approve_check.png", (2400, 800), (170, 170))
    d.text((2400, 920), "BANK ACCOUNT CREDITED", fill=P_GREEN, font=font_actor, anchor="mm")
    d.text((2400, 960), "Instant 24/7 Transfer", fill=P_MUTED, font=font_small, anchor="mm")
    d.line([(2020, 960), (2020, 1180), (2280, 1180)], fill=(0, 0, 0), width=4)
    draw_pil_arrow(d, (2270, 1180), (2300, 1180), "FAILED", (2150, 1140), head_len=18)
    paste_icon(bg, "scratch/extracted_elements/icon_reject_x.png", (2400, 1180), (160, 160))
    d.text((2400, 1300), "RESTORE WALLET", fill=P_RED, font=font_actor, anchor="mm")
    d.text((2400, 1340), "Pessimistic Lock & Rollback", fill=P_MUTED, font=font_small, anchor="mm")
    bg.save(f"{SLIDES_DIR}/slide_15.png")
    print("Slide 15 (Flow 5 Diagram) generated.")

# ==============================================================================
# SLIDE 16: FLOW 5 TEST MATRIX
# ==============================================================================
def build_slide_16(prs):
    succ = [
        ("TC_PAY_01: Quyết toán Thưởng mốc doanh số tháng (FR-09)", "Hệ thống quét tổng doanh số ròng qua Escrow, tính theo công thức bậc thang.\nTính đúng số tiền thưởng (ví dụ: 52.4tr * 2% = 1.048.000 ₫); Diễn giải rõ ràng."),
        ("TC_PAY_02: KOL tạo lệnh rút tiền & Khấu trừ thuế TNCN (FR-24)", "Lệnh rút tạo thành công ở PENDING_APPROVAL; Thực nhận 4.5tr (sau thuế 10% 500k).\nPessimistic Lock: Khóa số dư khả dụng ngay lập tức chống rút trùng."),
        ("TC_PAY_03: System Admin độc quyền duyệt Payout Napas (FR-24)", "Gọi API Napas 24/7 Instant Transfer; Tiền về tài khoản ngân hàng tức thì.\nLệnh chuyển COMPLETED; Xóa số dư phong tỏa vĩnh viễn khỏi ví."),
        ("TC_PAY_04: Cơ chế Rollback an toàn khi API Ngân hàng lỗi (FR-24)", "TUYỆT ĐỐI KHÔNG NUỐT LỖI: Ghi log chi tiết mã lỗi từ phía ngân hàng.\nDatabase Rollback: Mở khóa và hoàn trả 100% số dư phong tỏa về ví khả dụng.")
    ]
    fail = [
        ("TC_PAY_01: Đơn bị hoàn tiền hoặc KOL chưa đạt mốc tối thiểu", "Đơn đổi trả: Tự động trừ khỏi doanh số ròng trước khi áp dụng công thức.\nChưa đạt mốc: Ghi nhận 0đ thưởng tháng, giữ nguyên hoa hồng đơn."),
        ("TC_PAY_02: Số dư không đủ hoặc Rút dưới hạn mức tối thiểu", "Không đủ số dư: Form chặn submit, báo lỗi 'Số dư khả dụng không đủ'.\nRút dưới hạn mức (<100.000 ₫): Báo lỗi 'Số tiền rút tối thiểu 100.000 ₫'."),
        ("TC_PAY_03: User không phải System Admin hoặc STK sai lệch", "Không phải Super Admin: RBAC Guard chặn ngay lập tức với HTTP 403 Forbidden.\nAdmin từ chối lệnh: Bắt buộc nhập lý do, hoàn trả số dư về ví khả dụng."),
        ("TC_PAY_04: Lỗi Timeout mạng ngân hàng chưa rõ kết quả", "Lệnh đưa vào trạng thái PROCESSING_VERIFY; Cronjob đối soát 5 phút/lần.\nChỉ hoàn tiền sau khi đối soát đối ứng với Napas, bảo toàn ngân quỹ sàn.")
    ]
    build_test_matrix_slide(prs, 5, "MONTHLY BONUS & NAPAS PAYOUT AUDIT", succ, fail, 16)

# ==============================================================================
# SLIDE 17: THANK YOU & Q&A
# ==============================================================================
def build_slide_17(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "scratch/slide_assets/bg_thanks.png")

    add_editable_text(slide, 200, 560, 2600, 90, "THANK YOU FOR LISTENING", font_size_pt=46, bold=True, color=C_INK, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 200, 680, 2600, 120, "Q & A", font_size_pt=64, bold=True, color=C_BRAND_STRONG, align=PP_ALIGN.CENTER)
    add_editable_text(slide, 200, 830, 2600, 60, "Sales Collaborator & Affiliate Network Management System (SCANMS)", font_size_pt=20, bold=True, color=C_MUTED, align=PP_ALIGN.CENTER)

    bg = Image.open("scratch/slide_assets/bg_thanks.png").convert("RGB")
    d = ImageDraw.Draw(bg)
    d.text((W_PIXELS // 2, 600), "THANK YOU FOR LISTENING", fill=P_INK, font=font_title, anchor="mm")
    d.text((W_PIXELS // 2, 730), "Q & A", fill=P_BRAND_STRONG, font=font_hero, anchor="mm")
    d.text((W_PIXELS // 2, 850), "Sales Collaborator & Affiliate Network Management System (SCANMS)", fill=P_MUTED, font=font_actor, anchor="mm")
    bg.save(f"{SLIDES_DIR}/slide_17.png")
    print("Slide 17 (Q&A) generated.")

# ==============================================================================
# COMPILE HIGH-RES PDF WITH PYMUPDF (17 SLIDES)
# ==============================================================================
def compile_pdf():
    pdf_doc = fitz.open()
    for i in range(1, 18):
        img_path = f"{SLIDES_DIR}/slide_{i:02d}.png"
        img_doc = fitz.open(img_path)
        pdf_bytes = img_doc.convert_to_pdf()
        img_doc.close()
        slide_pdf = fitz.open("pdf", pdf_bytes)
        pdf_doc.insert_pdf(slide_pdf)
        slide_pdf.close()
    pdf_doc.save(PDF_OUT)
    pdf_doc.close()
    print(f"Master PDF compiled successfully (17 slides): {PDF_OUT}")

# ==============================================================================
# MAIN EXECUTION
# ==============================================================================
def main():
    prs = Presentation()
    prs.slide_width = Inches(W_INCHES)
    prs.slide_height = Inches(H_INCHES)

    print("Building all 17 master slides...")
    build_slide_01(prs)
    build_slide_02(prs)
    build_slide_03(prs)
    build_slide_04(prs)
    build_slide_05(prs)
    build_slide_06(prs)
    build_slide_07(prs) # Flow 1 Diagram
    build_slide_08(prs) # Flow 1 Test Matrix
    build_slide_09(prs) # Flow 2 Diagram
    build_slide_10(prs) # Flow 2 Test Matrix
    build_slide_11(prs) # Flow 3 Diagram
    build_slide_12(prs) # Flow 3 Test Matrix
    build_slide_13(prs) # Flow 4 Diagram
    build_slide_14(prs) # Flow 4 Test Matrix
    build_slide_15(prs) # Flow 5 Diagram
    build_slide_16(prs) # Flow 5 Test Matrix
    build_slide_17(prs) # Q&A / Thank You

    prs.save(PPTX_OUT)
    print(f"Master PPTX saved successfully: {PPTX_OUT}")

    compile_pdf()
    print("ALL 17 SLIDES COMPILED AND READY!")

if __name__ == "__main__":
    main()
