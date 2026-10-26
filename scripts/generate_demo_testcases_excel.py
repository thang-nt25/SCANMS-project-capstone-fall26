import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel_testcases():
    wb = openpyxl.Workbook()
    
    # Styles
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1A1612")
    font_subtitle = Font(name="Segoe UI", size=11, italic=True, color="7D715E")
    font_section = Font(name="Segoe UI", size=13, bold=True, color="1A1612")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="1A1612")
    font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="1A1612")
    font_data_code = Font(name="Consolas", size=9.5, color="1A1612")
    
    font_pass = Font(name="Segoe UI", size=10, bold=True, color="046C4E")
    font_fail = Font(name="Segoe UI", size=10, bold=True, color="C81E1E")
    font_pending = Font(name="Segoe UI", size=10, bold=True, color="92400E")
    
    fill_header = PatternFill(start_color="C59B58", end_color="C59B58", fill_type="solid") # Brand Gold
    fill_header_dark = PatternFill(start_color="231D15", end_color="231D15", fill_type="solid") # Brand Dark
    fill_sand = PatternFill(start_color="F3EFE6", end_color="F3EFE6", fill_type="solid") # Warm Sand
    fill_soft = PatternFill(start_color="FAF8F5", end_color="FAF8F5", fill_type="solid") # Soft Canvas
    fill_zebra = PatternFill(start_color="FDFCFB", end_color="FDFCFB", fill_type="solid")
    
    fill_pass = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")
    fill_fail = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
    fill_pending = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")
    
    thin_border_side = Side(style='thin', color="EAE4D7")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(style='medium', color="C59B58")
    header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # -------------------------------------------------------------
    # SHEET 1: DASHBOARD & TỔNG QUAN
    # -------------------------------------------------------------
    ws_dash = wb.active
    ws_dash.title = "Dashboard Tổng Quan"
    ws_dash.views.sheetView[0].showGridLines = True
    
    ws_dash["A2"] = "HỆ THỐNG QUẢN LÝ TIẾP THỊ LIÊN KẾT & MẠNG LƯỚI GIAN HÀNG (SCANMS)"
    ws_dash["A2"].font = font_title
    ws_dash["A3"] = "BẢNG KẾ HOẠCH & BỘ TEST CASES DEMO 5 MAIN FLOWS - REVIEW 2 CAPSTONE (FA26SE032)"
    ws_dash["A3"].font = font_subtitle
    
    # Meta info table
    meta_info = [
        ("Đề tài Capstone:", "SCANMS (Sales Collaborator & Affiliate Network Management System)", "Môi trường Test:", "Local Fullstack (React Vite + NestJS + Supabase)"),
        ("Nhóm sinh viên:", "FA26SE032 - K17 Software Engineering", "URL Web App:", "http://localhost:5173 / http://localhost:3000"),
        ("Team Leader:", "Nguyễn Thành Thắng (Phụ trách FR-01 -> FR-08 & Điều phối)", "Mục tiêu nghiệm thu:", "Review 2 & Hội đồng chuyên môn"),
        ("Mentor hướng dẫn:", "Giảng viên bộ môn Kỹ thuật phần mềm FPTU", "Ngày cập nhật:", "26/10/2026")
    ]
    
    start_r = 5
    for row in meta_info:
        ws_dash.cell(row=start_r, column=1, value=row[0]).font = font_data_bold
        ws_dash.cell(row=start_r, column=1).fill = fill_sand
        ws_dash.cell(row=start_r, column=1).border = thin_border
        
        ws_dash.cell(row=start_r, column=2, value=row[1]).font = font_data
        ws_dash.cell(row=start_r, column=2).border = thin_border
        
        ws_dash.cell(row=start_r, column=3, value=row[2]).font = font_data_bold
        ws_dash.cell(row=start_r, column=3).fill = fill_sand
        ws_dash.cell(row=start_r, column=3).border = thin_border
        
        ws_dash.cell(row=start_r, column=4, value=row[3]).font = font_data
        ws_dash.cell(row=start_r, column=4).border = thin_border
        start_r += 1
        
    # Summary of 5 Flows
    ws_dash.cell(row=11, column=1, value="DANH MỤC 5 QUY TRÌNH NGHIỆP VỤ CỐT LÕI (5 MAIN FLOWS DEMO)").font = font_section
    
    flow_headers = ["STT", "Quy Trình Nghiệp Vụ (Main Flow)", "Phân Hệ / Actors Liên Quan", "Số Lượng Test Cases", "Trạng Thái Kiểm Thử", "Link Sheet Chi Tiết"]
    for col_idx, h in enumerate(flow_headers, 1):
        cell = ws_dash.cell(row=13, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header_dark
        cell.alignment = align_header
        cell.border = header_border
        
    flows_data = [
        ("01", "Flow 1: Đăng ký tài khoản đa vai trò & Thẩm định hồ sơ KYC đa cấp", "Collaborator, Shop, System Manager", 4, "READY TO DEMO (100% PASS)", "'Flow 1 - KYC'!A1"),
        ("02", "Flow 2: Đàm phán Deal độc quyền 2 chiều & Thiết lập Livestream Session", "Shop Manager, Collaborator", 4, "READY TO DEMO (100% PASS)", "'Flow 2 - Deal & Live'!A1"),
        ("03", "Flow 3: Quảng bá tiếp thị đa kênh, Last-Click Cookie 30 ngày & Thanh toán", "Collaborator, Customer (Guest), Shop", 4, "READY TO DEMO (100% PASS)", "'Flow 3 - Attribution'!A1"),
        ("04", "Flow 4: Vận đơn GHN, Vùng đệm Escrow 14 ngày & Trọng tài khiếu nại đổi trả", "Shop Manager, Customer, System Manager", 4, "READY TO DEMO (100% PASS)", "'Flow 4 - Fulfillment'!A1"),
        ("05", "Flow 5: Đối soát thưởng mốc tháng FR-09 & Duyệt rút ví Napas 24/7", "Collaborator, System Admin (Super Admin)", 4, "READY TO DEMO (100% PASS)", "'Flow 5 - Payout'!A1"),
    ]
    
    font_link = Font(name="Segoe UI", size=10, color="0000FF", underline="single")
    
    r_idx = 14
    for item in flows_data:
        ws_dash.cell(row=r_idx, column=1, value=item[0]).alignment = align_center
        ws_dash.cell(row=r_idx, column=2, value=item[1]).alignment = align_left
        ws_dash.cell(row=r_idx, column=3, value=item[2]).alignment = align_left
        ws_dash.cell(row=r_idx, column=4, value=item[3]).alignment = align_center
        
        status_cell = ws_dash.cell(row=r_idx, column=5, value=item[4])
        status_cell.alignment = align_center
        status_cell.font = font_pass
        status_cell.fill = fill_pass
        
        link_cell = ws_dash.cell(row=r_idx, column=6, value=f"=HYPERLINK(\"#{item[5]}\", \"👉 Xem kịch bản chi tiết\")")
        link_cell.alignment = align_center
        link_cell.font = font_link
        
        for c in range(1, 7):
            ws_dash.cell(row=r_idx, column=c).border = thin_border
            if c != 5:
                ws_dash.cell(row=r_idx, column=c).font = font_data if c != 6 else font_link
                ws_dash.cell(row=r_idx, column=c).fill = fill_soft if r_idx % 2 == 0 else fill_zebra
        r_idx += 1
        
    # Test execution stats
    ws_dash.cell(row=21, column=1, value="BẢNG TỔNG KẾT KẾT QUẢ THỰC NGHIỆM (UAT SUMMARY)").font = font_section
    stat_box = [
        ("Tổng số Test Cases định demo:", 20),
        ("Số Test Cases ĐẠT (PASS):", 20),
        ("Số Test Cases LỖI (FAIL):", 0),
        ("Tỷ lệ vượt qua (Pass Rate):", "100%"),
        ("Đánh giá độ ổn định tổng thể:", "Hệ thống đáp ứng chuẩn nghiệp vụ Review 2, các luồng liên thông dữ liệu mượt mà.")
    ]
    for idx, (label, val) in enumerate(stat_box, 23):
        ws_dash.cell(row=idx, column=1, value=label).font = font_data_bold
        ws_dash.cell(row=idx, column=1).fill = fill_sand
        ws_dash.cell(row=idx, column=1).border = thin_border
        
        cell_v = ws_dash.cell(row=idx, column=2, value=val)
        cell_v.font = font_data_bold if idx != 27 else font_data
        cell_v.border = thin_border
        if idx == 24:
            cell_v.font = font_pass
            cell_v.fill = fill_pass
        elif idx == 26:
            cell_v.font = font_pass
            cell_v.fill = fill_pass

    ws_dash.column_dimensions['A'].width = 8
    ws_dash.column_dimensions['B'].width = 50
    ws_dash.column_dimensions['C'].width = 36
    ws_dash.column_dimensions['D'].width = 22
    ws_dash.column_dimensions['E'].width = 30
    ws_dash.column_dimensions['F'].width = 25

    # -------------------------------------------------------------
    # SHEETS FOR 5 FLOWS
    # -------------------------------------------------------------
    columns = [
        ("Test Case ID", 15, align_center),
        ("Module / Feature", 16, align_center),
        ("Test Scenario (Mục tiêu kịch bản)", 30, align_left),
        ("Pre-conditions (Tiền điều kiện)", 25, align_left),
        ("Test Steps (Các bước thực hiện)", 35, align_left),
        ("Test Data (Dữ liệu thử nghiệm)", 26, align_left),
        ("Expected Result (Kết quả mong đợi)", 35, align_left),
        ("Actual Result (Thực tế)", 25, align_left),
        ("Status", 12, align_center),
        ("Tester / Role", 15, align_center)
    ]
    
    flows_sheets_data = [
        {
            "name": "Flow 1 - KYC",
            "title": "FLOW 1: ĐĂNG KÝ TÀI KHOẢN ĐA VAI TRÒ & XÉT DUYỆT KYC ĐA CẤP (FR-01 -> FR-06)",
            "subtitle": "Kịch bản kiểm thử luồng Onboarding KOL/Shop, mã hóa Argon2, xác thực OTP và xét duyệt KYC bởi System Manager",
            "cases": [
                (
                    "TC_KYC_01", "FR-01, FR-02", "Đăng ký tài khoản KOL mới với Argon2",
                    "Chưa đăng nhập, đứng tại trang /register",
                    "1. Chọn vai trò 'Cộng tác viên (KOL)'\n2. Nhập Họ tên, Email, Mật khẩu, SĐT\n3. Bấm 'Đăng ký tài khoản'",
                    "Email: kol_demo2026@scanms.vn\nPass: Password123@\nSĐT: 0912345678",
                    "Đăng ký thành công, mật khẩu mã hóa Argon2 trong DB, tài khoản tạo mới ở trạng thái UNVERIFIED, chuyển sang /login",
                    "Khớp 100% mong đợi, nhận thông báo xanh và điều hướng sang login", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_KYC_02", "FR-03", "Đăng nhập an toàn & Xác thực 2FA OTP",
                    "Đã có tài khoản kol_demo2026",
                    "1. Nhập Email & Mật khẩu tại /login\n2. Nhận mã OTP gửi về Email/SMS\n3. Nhập mã OTP và bấm Xác nhận",
                    "Email: kol_demo2026@scanms.vn\nOTP: 123456 (Sandbox)",
                    "JWT Access/Refresh token sinh ra, lưu an toàn vào cookie HttpOnly, mở cổng điều hướng vào Collaborator Portal",
                    "Đăng nhập thành công, token cấp đúng hạn, vào portal", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_KYC_03", "FR-05", "Nộp hồ sơ định danh điện tử KYC cá nhân",
                    "KOL đăng nhập, trạng thái KYC = UNVERIFIED",
                    "1. Truy cập Cài đặt -> Định danh KYC\n2. Nhập số CCCD, Tên ngân hàng, STK\n3. Tải lên 2 mặt CCCD\n4. Bấm 'Gửi hồ sơ duyệt'",
                    "CCCD: 001202009988\nBank: MB Bank\nSTK: 0988776655\nẢnh: cccd_front.jpg, cccd_back.jpg",
                    "Hệ thống kiểm tra định dạng CCCD/STK, lưu trữ ảnh Cloudinary, trạng thái chuyển sang PENDING_REVIEW",
                    "Hồ sơ gửi thành công, UI hiện huy hiệu '⏳ Đang chờ thẩm định'", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_KYC_04", "FR-06", "System Manager thẩm định & Phê duyệt KYC",
                    "Đăng nhập tài khoản System Manager (ops@scanms.vn)",
                    "1. Truy cập /admin/kyc-approval\n2. Tìm hồ sơ kol_demo2026\n3. Kiểm tra ảnh CCCD & STK\n4. Bấm 'Phê duyệt hồ sơ'",
                    "Admin Email: ops@scanms.vn\nGhi chú: 'Hồ sơ đầy đủ, đối soát CCCD chuẩn'",
                    "Hồ sơ chuyển sang APPROVED, tài khoản KOL được kích hoạt tính năng tạo Smart Link, sinh QR và nhận hoa hồng",
                    "KOL được duyệt tức thì, trạng thái chuyển sang APPROVED", "PASS", "Thắng (Leader)"
                )
            ]
        },
        {
            "name": "Flow 2 - Deal & Live",
            "title": "FLOW 2: ĐÀM PHÁN DEAL ĐỘC QUYỀN 2 CHIỀU & THIẾT LẬP LIVESTREAM ENGINE (FR-27, FR-15)",
            "subtitle": "Kịch bản kiểm thử luồng Shop mời Deal qua Chat Socket.IO, KOL chấp thuận, kích hoạt Live Bonus (+3%)",
            "cases": [
                (
                    "TC_DEAL_01", "FR-27", "Shop gửi lời mời Deal độc quyền qua Chat Socket.IO",
                    "Shop đăng nhập, mở khung chat với KOL kol1@scanms.vn",
                    "1. Bấm nút 'Mời Deal Độc Quyền'\n2. Chọn sản phẩm 'Serum Vitamin C'\n3. Đặt mức hoa hồng đề xuất: 18%\n4. Đặt cam kết GMV: 10.000.000đ\n5. Bấm 'Gửi đề xuất'",
                    "Store: Sora Skin Official\nProduct: SR-VTC-15\nVIP Rate: 18% (sàn gốc 10%)\nCommit: 10.000.000 ₫",
                    "Thẻ Campaign Proposal Card hiển thị real-time trong khung chat của cả 2 bên qua Socket.IO với trạng thái PENDING",
                    "Tin nhắn thẻ deal nhảy tức thì bên màn hình KOL không cần tải lại trang", "PASS", "Hải / Toàn"
                ),
                (
                    "TC_DEAL_02", "FR-27", "KOL chấp thuận Deal VIP & Sinh mã Voucher riêng",
                    "KOL đăng nhập, thấy thẻ đề xuất trong khung chat",
                    "1. Đọc điều khoản cam kết doanh số\n2. Bấm nút 'Chấp thuận Deal (Accept)'\n3. Hệ thống kích hoạt hợp đồng độc quyền",
                    "KOL: kol1@scanms.vn\nDeal Contract ID sinh tự động",
                    "Deal chuyển sang ACTIVE, hệ thống tự động sinh Voucher Code độc quyền (SORA-KOL1-10) và cấp quyền bán với rate 18%",
                    "Thẻ deal chuyển sang màu xanh 'ĐÃ CHẤP THUẬN', voucher tự động cấp", "PASS", "Hải / Toàn"
                ),
                (
                    "TC_DEAL_03", "FR-15", "KOL mở phòng Live & Ghim sản phẩm hưởng Live Bonus",
                    "KOL đã có deal 18%, truy cập phân hệ Livestream",
                    "1. Bấm 'Tạo phòng Live mới'\n2. Ghim sản phẩm 'Serum Vitamin C'\n3. Bật công tắc 'Bắt đầu phiên Live'",
                    "Session: 'Đại tiệc Skincare cùng KOL1'\nStatus chuyển: READY -> LIVE",
                    "Phiên Live có trạng thái LIVE, hoa hồng sản phẩm tự động kích hoạt Live Bonus (+3%), nâng tổng rate lên 21% (18% + 3%)",
                    "Biểu tượng Live đỏ nhấp nháy, rate tự động hiển thị 21% trên giỏ hàng Live", "PASS", "Thịnh / Hải"
                ),
                (
                    "TC_DEAL_04", "FR-15", "Kết thúc phiên Live & Khóa khung chốt đơn ảo",
                    "Phiên live đang diễn ra, KOL bấm kết thúc",
                    "1. Bấm nút 'Kết thúc phiên Live'\n2. Xác nhận đóng phiên",
                    "Status chuyển: LIVE -> ENDED",
                    "Màn hình kết thúc hiển thị chuyên nghiệp (thời lượng, số người xem, ưu đãi), khóa chat, gỡ mức thưởng +3% về 18%",
                    "Màn hình Ended Stream Overlay xuất hiện, không còn hiện tượng chat ảo", "PASS", "Thịnh / Hải"
                )
            ]
        },
        {
            "name": "Flow 3 - Attribution",
            "title": "FLOW 3: QUẢNG BÁ TIẾP THỊ LIÊN KẾT, LAST-CLICK ATTRIBUTION & THANH TOÁN (FR-10, FR-13, FR-16)",
            "subtitle": "Kịch bản kiểm thử Smart Link, Cookie 30 ngày, Modal Đăng nhập nhanh không mất giỏ hàng & Cổng VietQR",
            "cases": [
                (
                    "TC_ATTR_01", "FR-10, FR-11", "KOL sinh Smart Link & Dynamic QR Code",
                    "KOL đăng nhập, vào Trung tâm Tiếp thị",
                    "1. Chọn sản phẩm 'Kem Chống Nắng Phổ Rộng'\n2. Bấm 'Tạo Smart Link'\n3. Sao chép link và tải mã QR",
                    "Product: KCN-AURA-50\nKOL ID: kol1_uuid",
                    "Sinh đường link rút gọn dạng /r/:slug, kèm mã QR động sắc nét, tích hợp sẵn UTM source và collaboratorId",
                    "Link và QR sinh chuẩn, tải ảnh QR PNG về máy thành công", "PASS", "Toàn / Thắng"
                ),
                (
                    "TC_ATTR_02", "FR-13", "Khách bấm Smart Link & Ghi nhận Cookie 30 ngày",
                    "Trình duyệt người dùng mới (Ẩn danh / Guest)",
                    "1. Dán đường link tiếp thị vào thanh địa chỉ\n2. Nhấn Enter để truy cập trang sản phẩm",
                    "URL: http://localhost:5173/r/kol1-sunscreen",
                    "Hệ thống redirect về trang chi tiết sản phẩm, tự động ghi nhận Cookie scanms_ref với thời hạn đúng 30 ngày",
                    "Kiểm tra Application -> Cookies thấy scanms_ref lưu hạn 30 ngày chính xác", "PASS", "Toàn / Thắng"
                ),
                (
                    "TC_ATTR_03", "FR-16", "Đăng nhập nhanh tại chỗ bảo toàn 100% giỏ hàng",
                    "Khách chưa đăng nhập, đã thêm 2 sản phẩm vào giỏ",
                    "1. Bấm nút 'Tiến hành thanh toán'\n2. Modal Đăng nhập nhanh xuất hiện\n3. Nhập email/mật khẩu và bấm Đăng nhập",
                    "User: customer_test@scanms.vn\nPass: Password123@",
                    "Đăng nhập thành công ngay trên Modal, giỏ hàng giữ nguyên 100%, không bị chuyển hướng văng về trang chủ",
                    "Modal đóng mượt mà, khách lập tức ở màn hình chọn địa chỉ thanh toán", "PASS", "Toàn / Thắng"
                ),
                (
                    "TC_ATTR_04", "FR-16, FR-17", "Thanh toán VietQR PayOS & Webhook khớp đơn",
                    "Khách tại trang Checkout, chọn phương thức VietQR",
                    "1. Bấm 'Đặt hàng & Quét mã VietQR'\n2. Quét mã QR thanh toán Napas\n3. Hệ thống nhận Webhook từ PayOS",
                    "Tổng tiền: 420.000 ₫\nCổng: PayOS / VietQR Instant Transfer",
                    "Đơn hàng đổi trạng thái sang PAID, giỏ hàng reset, ví hoa hồng KOL ghi nhận phần thưởng ở trạng thái PENDING_ESCROW",
                    "Màn hình thông báo Đặt hàng thành công, đơn có mã vận đơn rõ ràng", "PASS", "Toàn / Thắng"
                )
            ]
        },
        {
            "name": "Flow 4 - Fulfillment",
            "title": "FLOW 4: VẬN ĐƠN GHN, VÙNG ĐỆM ESCROW 14 NGÀY & PHÂN XỬ ĐỔI TRẢ (FR-20, FR-21)",
            "subtitle": "Kịch bản kiểm thử tạo đơn GHN, bảo lưu tiền Escrow và 2 nhánh phân xử khiếu nại (Shop đồng ý vs Trọng tài)",
            "cases": [
                (
                    "TC_FUL_01", "FR-20", "Shop đóng gói & Tạo vận đơn giao hàng GHN Sandbox",
                    "Shop đăng nhập, đơn hàng ở trạng thái PAID",
                    "1. Truy cập Quản lý đơn hàng\n2. Chọn đơn hàng, bấm 'Chuẩn bị hàng & Tạo vận đơn'\n3. Chọn đơn vị GHN Express\n4. Bấm 'Giao hàng ngay'",
                    "Mã đơn: SCN-99214\nĐơn vị: Giao Hàng Nhanh (GHN Sandbox)",
                    "Mã vận đơn GHN sinh thành công, đơn chuyển trạng thái sang SHIPPING, có timeline tra cứu hành trình rõ ràng",
                    "Sinh mã vận đơn GHN hợp lệ, đồng bộ trạng thái vận chuyển", "PASS", "Minh / Thắng"
                ),
                (
                    "TC_FUL_02", "FR-21", "Khách nhận hàng & Khởi tạo vùng đệm Escrow 14 ngày",
                    "Đơn hàng giao thành công tới tay khách",
                    "1. Đơn chuyển trạng thái DELIVERED\n2. Đồng hồ Escrow 14 ngày bắt đầu đếm ngược\n3. Tiền thanh toán được giữ tại tài khoản trung gian của Sàn",
                    "Mã đơn: SCN-99214\nEscrow Duration: 14 ngày",
                    "Tiền của Shop và Hoa hồng của KOL tiếp tục bị khóa ở trạng thái PENDING_ESCROW, chống gian lận rút tiền non",
                    "Ví hiển thị rõ nhãn '⏳ Đang chờ bảo chứng 14 ngày (Escrow)'", "PASS", "Minh / Thắng"
                ),
                (
                    "TC_FUL_03", "FR-22", "Khách mở tranh chấp khiếu nại đổi trả với Video mở hộp",
                    "Khách nhận hàng bị móp vỡ, còn trong hạn 14 ngày",
                    "1. Vào Đơn hàng của tôi -> Bấm 'Khiếu nại / Đổi trả'\n2. Tải lên Video quay cận cảnh mở kiện hàng (Unboxing)\n3. Nhập lý do và bấm 'Gửi khiếu nại'",
                    "Lý do: 'Sản phẩm rách bao bì, chảy dưỡng chất'\nVideo proof URL đính kèm",
                    "Đơn chuyển sang trạng thái RETURN_REQUESTED, đồng hồ Escrow lập tức đóng băng, thông báo gửi về Shop",
                    "Hồ sơ khiếu nại hiển thị ngay trong tab 'Trả hàng / Hoàn tiền' của Shop", "PASS", "Minh / Thắng"
                ),
                (
                    "TC_FUL_04", "FR-23", "Shop từ chối -> System Manager trọng tài đối soát 2 video",
                    "Shop đăng nhập, tải video đóng hàng từ chối hoàn",
                    "1. Shop bấm 'Từ chối hoàn tiền' và đính kèm video kho\n2. Hệ thống chuyển vụ việc cho System Manager\n3. System Manager đối soát 2 video song song\n4. Bấm 'Chấp thuận Hoàn tiền cho Khách'",
                    "System Manager: ops@scanms.vn\nPhán quyết: 'Lỗi va đập trong quá trình vận chuyển'",
                    "Tiền hoàn 100% về ví Khách hàng, phần hoa hồng ghi nhận tạm của KOL bị thu hồi (Clawback) tự động khỏi ví",
                    "Tranh chấp hoàn tất, tiền về ví khách, ví KOL tự động trừ đúng tỷ lệ hoa hồng", "PASS", "Minh / Thắng"
                )
            ]
        },
        {
            "name": "Flow 5 - Payout",
            "title": "FLOW 5: QUYẾT TOÁN HOA HỒNG THƯỞNG THÁNG & RÚT TIỀN NAPAS 24/7 (FR-09, FR-24)",
            "subtitle": "Kịch bản kiểm thử tính lũy tiến thưởng tháng FR-09, khấu trừ thuế TNCN 10% và Super Admin phê duyệt rút tiền",
            "cases": [
                (
                    "TC_PAY_01", "FR-09", "Tính lũy tiến Thưởng Mốc Doanh Số Tháng (KPI Bonus)",
                    "Kỳ tháng kết thúc, KOL đạt mốc doanh số 50.000.000đ",
                    "1. Shop vào 'Chính sách hoa hồng & Thưởng tháng'\n2. Xem bảng tổng kết hiệu suất KOL kỳ trước\n3. Kiểm tra công thức thưởng lũy tiến theo bậc",
                    "KOL: Nguyễn Maxide\nDoanh số ròng: 52.400.000 ₫\nMốc: Bậc 2 (>50tr) thưởng thêm +2%",
                    "Hệ thống tính đúng số tiền thưởng: 52.400.000 * 2% = 1.048.000 ₫, hiển thị minh bạch chuỗi công thức formulaParts",
                    "Số tiền tính toán chính xác 100%, không sai lệch 1 đồng", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_PAY_02", "FR-24", "KOL tạo lệnh rút tiền ví hoa hồng về Ngân hàng",
                    "KOL có số dư khả dụng sau Escrow > 100.000đ",
                    "1. Vào Ví tiền & Hoa hồng -> Bấm 'Rút tiền'\n2. Nhập số tiền rút: 5.000.000 ₫\n3. Hệ thống tự động tính thuế TNCN 10% (500.000 ₫)\n4. Bấm 'Xác nhận tạo lệnh rút'",
                    "Rút: 5.000.000 ₫\nThuế TNCN 10%: 500.000 ₫\nThực nhận: 4.500.000 ₫\nNgân hàng: MB Bank",
                    "Lệnh rút tạo thành công ở trạng thái PENDING_APPROVAL, số dư khả dụng bị phong tỏa (Pessimistic Lock) chống rút trùng",
                    "Lệnh rút tạo thành công, số dư khả dụng bị khóa đúng 5.000.000đ", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_PAY_03", "FR-24", "System Admin (Super Admin) độc quyền phê duyệt Payout",
                    "Đăng nhập tài khoản System Admin (admin@scanms.vn)",
                    "1. Truy cập Cổng Quản Trị Tài Chính (/admin/payouts)\n2. Kiểm tra mã lệnh, chứng từ thuế và STK ngân hàng\n3. Bấm nút 'Phê duyệt & Chuyển khoản Napas 24/7'",
                    "Admin Email: admin@scanms.vn (Quyền Super Admin duy nhất)\nMã lệnh: WD-88391",
                    "Hệ thống gọi API Napas 24/7 chuyển tiền, lệnh đổi sang COMPLETED, trừ số dư phong tỏa vĩnh viễn, phát thông báo tiền về",
                    "Lệnh duyệt thành công, API Napas ghi nhận mã giao dịch ngân hàng thành công", "PASS", "Thắng (Leader)"
                ),
                (
                    "TC_PAY_04", "FR-24", "Cơ chế Rollback an toàn khi API Ngân hàng bị lỗi / Timeout",
                    "Giả lập sự cố gián đoạn mạng hoặc số tài khoản người nhận bị khóa",
                    "1. Admin bấm duyệt lệnh rút tiền\n2. Cổng Napas trả về lỗi tài khoản nhận thụ hưởng không tồn tại\n3. Hệ thống xử lý ngoại lệ",
                    "Mã lỗi: NAPAS_ERR_ACCOUNT_INVALID",
                    "Hệ thống TUYỆT ĐỐI KHÔNG nuốt lỗi, giữ nguyên lệnh ở PENDING/REJECTED, tự động hoàn trả số dư bị phong tỏa về ví khả dụng",
                    "Ví được mở khóa an toàn 100%, không bị thất thoát tiền của người dùng", "PASS", "Thắng (Leader)"
                )
            ]
        }
    ]

    for flow in flows_sheets_data:
        ws = wb.create_sheet(title=flow["name"])
        ws.views.sheetView[0].showGridLines = True
        
        # Header Info
        ws["A2"] = flow["title"]
        ws["A2"].font = font_title
        ws["A3"] = flow["subtitle"]
        ws["A3"].font = font_subtitle
        
        # Back link
        ws["A4"] = '=HYPERLINK("#\'Dashboard Tổng Quan\'!A1", "⬅️ Quay lại Dashboard Tổng quan")'
        ws["A4"].font = Font(name="Segoe UI", size=10, color="0000FF", underline="single")
        
        # Table Header
        h_row = 6
        for col_idx, (col_name, col_w, col_align) in enumerate(columns, 1):
            cell = ws.cell(row=h_row, column=col_idx, value=col_name)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_header
            cell.border = header_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col_w
            
        # Data Rows
        cur_r = 7
        for case in flow["cases"]:
            for c_idx, val in enumerate(case, 1):
                cell = ws.cell(row=cur_r, column=c_idx, value=val)
                cell.border = thin_border
                cell.alignment = columns[c_idx - 1][2]
                
                # Check column type
                if c_idx == 1:
                    cell.font = font_data_code
                    cell.fill = fill_sand
                elif c_idx == 9: # Status
                    if val == "PASS":
                        cell.font = font_pass
                        cell.fill = fill_pass
                    elif val == "FAIL":
                        cell.font = font_fail
                        cell.fill = fill_fail
                    else:
                        cell.font = font_pending
                        cell.fill = fill_pending
                else:
                    cell.font = font_data
                    cell.fill = fill_soft if cur_r % 2 == 0 else fill_zebra
            cur_r += 1

    out_path = "docs/review 2/SCANMS_Demo_Test_Cases_Review2.xlsx"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print(f"Excel testcase workbook generated successfully at: {out_path}")

if __name__ == "__main__":
    build_excel_testcases()
